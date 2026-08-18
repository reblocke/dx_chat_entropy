from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
import zipfile
from collections import defaultdict
from collections.abc import Mapping, Sequence
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .manifest import read_manifest
from .models import FeedbackRequest
from .runtime import validate_response_record

WORKBOOK_MANIFEST_SCHEMA_VERSION = "feedback-workbook-manifest-v1"
INCOMPLETE_SHEET_NAME = "INCOMPLETE"
_FIXED_WORKBOOK_TIME = datetime(2000, 1, 1)
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_SURFACE_DIRECTORIES = {
    "overall/general": "overall_gen",
    "overall/specific": "overall_spec",
    "differential/general": "diff_gen",
    "differential/specific": "diff_spec",
}


class WorkbookMaterializationError(ValueError):
    """Raised when validated records cannot satisfy the workbook contract."""


def _fsync_directory(directory: Path) -> None:
    try:
        descriptor = os.open(directory, os.O_RDONLY)
    except OSError:  # pragma: no cover - unsupported filesystems/platforms.
        return
    try:
        os.fsync(descriptor)
    except OSError:  # pragma: no cover - unsupported filesystems/platforms.
        pass
    finally:
        os.close(descriptor)


def _safe_run_path(run_dir: Path, relative_path: str) -> Path:
    path = Path(relative_path)
    if path.is_absolute() or ".." in path.parts:
        raise WorkbookMaterializationError(
            f"Run artifact path must be safe and relative: {relative_path!r}"
        )
    resolved = (run_dir / path).resolve()
    try:
        resolved.relative_to(run_dir.resolve())
    except ValueError as exc:
        raise WorkbookMaterializationError(
            f"Run artifact path escapes the run directory: {relative_path!r}"
        ) from exc
    return resolved


def _partial_path(path: Path) -> Path:
    return path.with_name(f"{path.stem}_PARTIAL{path.suffix}")


def expected_workbook_path(request: FeedbackRequest) -> str:
    """Return the frozen notebook-compatible destination for a request."""

    try:
        directory = _SURFACE_DIRECTORIES[request.surface]
    except KeyError as exc:
        raise WorkbookMaterializationError(
            f"{request.request_id}: unsupported surface {request.surface!r}"
        ) from exc
    if request.surface == "overall/general":
        filename = f"{request.target_diagnosis}_gen_overall.xlsx"
    elif request.surface == "overall/specific":
        filename = f"{request.target_diagnosis}_spec_overall.xlsx"
    else:
        if not request.comparison_diagnosis:
            raise WorkbookMaterializationError(
                f"{request.request_id}: differential request lacks comparison_diagnosis"
            )
        filename = f"{request.correct_diagnosis}_vs_{request.comparison_diagnosis}.xlsx"
    return (Path("workbooks") / directory / filename).as_posix()


def validate_request_destinations(request: FeedbackRequest) -> None:
    """Validate non-identity artifact fields against the frozen output contract."""

    if request.surface.startswith("overall/"):
        if request.comparison_diagnosis is not None:
            raise WorkbookMaterializationError(
                f"{request.request_id}: overall request has a comparison diagnosis"
            )
        if request.expected_response_type != "overall":
            raise WorkbookMaterializationError(
                f"{request.request_id}: overall request has the wrong response type"
            )
    elif request.surface.startswith("differential/"):
        if request.target_diagnosis != request.correct_diagnosis:
            raise WorkbookMaterializationError(
                f"{request.request_id}: differential target must be the correct diagnosis"
            )
        if request.comparison_diagnosis in {None, request.correct_diagnosis}:
            raise WorkbookMaterializationError(
                f"{request.request_id}: differential comparison diagnosis is invalid"
            )
        if request.expected_response_type != "differential":
            raise WorkbookMaterializationError(
                f"{request.request_id}: differential request has the wrong response type"
            )

    expected_response_path = f"responses/{request.request_id}.json"
    if request.response_path != expected_response_path:
        raise WorkbookMaterializationError(
            f"{request.request_id}: response_path differs from the request contract"
        )
    if request.workbook_path != expected_workbook_path(request):
        raise WorkbookMaterializationError(
            f"{request.request_id}: workbook_path differs from the legacy contract"
        )
    if request.sheet_name != request.category_key[:31]:
        raise WorkbookMaterializationError(
            f"{request.request_id}: sheet_name differs from the category contract"
        )
    if request.initial_status != "pending":
        raise WorkbookMaterializationError(
            f"{request.request_id}: initial_status must be 'pending'"
        )


def _payload_lists(
    request: FeedbackRequest, payload: Mapping[str, object]
) -> tuple[Sequence[object], Sequence[object]]:
    if request.surface.startswith("overall/"):
        keys = (
            "for_diagnosis_strongest_evidence",
            "against_diagnosis_strongest_evidence",
        )
    elif request.surface.startswith("differential/"):
        keys = ("diagnosisA_strongest_evidence", "diagnosisB_strongest_evidence")
    else:
        raise WorkbookMaterializationError(
            f"{request.request_id}: unsupported surface {request.surface!r}"
        )

    values: list[Sequence[object]] = []
    for key in keys:
        value = payload.get(key)
        if not isinstance(value, Sequence) or isinstance(value, (str, bytes, bytearray)):
            raise WorkbookMaterializationError(
                f"{request.request_id}: parsed_payload.{key} must be an array"
            )
        if len(value) != 5:
            raise WorkbookMaterializationError(
                f"{request.request_id}: parsed_payload.{key} must contain exactly 5 items"
            )
        values.append(value)
    return values[0], values[1]


def _item_text(item: object, field: str, *, request_id: str) -> str:
    if not isinstance(item, Mapping):
        raise WorkbookMaterializationError(f"{request_id}: evidence item must be an object")
    value = item.get(field)
    if not isinstance(value, str) or not value.strip():
        raise WorkbookMaterializationError(
            f"{request_id}: evidence item {field} must be a non-empty string"
        )
    return value


def _headers(request: FeedbackRequest) -> tuple[str, str, str, str]:
    if request.surface.startswith("overall/"):
        diagnosis = request.target_diagnosis
        return (
            f"Supports {diagnosis}",
            f"Rationale for {diagnosis}",
            f"Against {diagnosis}",
            f"Rationale Against {diagnosis}",
        )

    comparison = request.comparison_diagnosis
    if not comparison:
        raise WorkbookMaterializationError(
            f"{request.request_id}: differential request lacks comparison_diagnosis"
        )
    return (
        f"Supports {request.correct_diagnosis}",
        f"Rationale {request.correct_diagnosis}",
        f"Supports {comparison}",
        f"Rationale {comparison}",
    )


def _table_for(
    request: FeedbackRequest, record: Mapping[str, object]
) -> tuple[tuple[str, str, str, str], list[tuple[str, str, str, str]]]:
    payload = record.get("parsed_payload")
    if not isinstance(payload, Mapping):
        raise WorkbookMaterializationError(
            f"{request.request_id}: response record lacks parsed_payload"
        )
    first, second = _payload_lists(request, payload)
    rows = [
        (
            _item_text(first[index], "finding", request_id=request.request_id),
            _item_text(first[index], "explanation", request_id=request.request_id),
            _item_text(second[index], "finding", request_id=request.request_id),
            _item_text(second[index], "explanation", request_id=request.request_id),
        )
        for index in range(5)
    ]
    return _headers(request), rows


def _validate_sheet_contract(requests: Sequence[FeedbackRequest]) -> None:
    seen: set[str] = set()
    for request in requests:
        name = request.sheet_name
        if not name or len(name) > 31 or any(character in name for character in r"[]:*?/\\"):
            raise WorkbookMaterializationError(
                f"{request.request_id}: invalid Excel sheet name {name!r}"
            )
        if name == INCOMPLETE_SHEET_NAME:
            raise WorkbookMaterializationError(
                f"{request.request_id}: {INCOMPLETE_SHEET_NAME!r} is reserved"
            )
        if name in seen:
            raise WorkbookMaterializationError(f"Workbook has duplicate sheet name {name!r}")
        seen.add(name)


def _set_fixed_properties(workbook: Workbook) -> None:
    properties = workbook.properties
    properties.creator = "dx_chat_entropy"
    properties.lastModifiedBy = "dx_chat_entropy"
    properties.created = _FIXED_WORKBOOK_TIME
    properties.modified = _FIXED_WORKBOOK_TIME
    properties.title = "Clinical diagnostic-reasoning feedback"
    properties.subject = "Validated model-generated feedback"
    properties.description = "Generated deterministically from validated response records."
    properties.keywords = "feedback, model output"
    properties.category = "Research workflow artifact"
    properties.revision = "1"


def _append_table(
    worksheet: Worksheet,
    headers: tuple[str, str, str, str],
    rows: Sequence[tuple[str, str, str, str]],
) -> None:
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)


def _canonicalize_zip(source: Path, destination: Path) -> None:
    """Rewrite an XLSX ZIP with fixed metadata and member ordering."""

    with zipfile.ZipFile(source, "r") as archive:
        members = {info.filename: archive.read(info.filename) for info in archive.infolist()}

    core_path = "docProps/core.xml"
    if core_path in members:
        members[core_path] = re.sub(
            rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
            rb"\g<1>2000-01-01T00:00:00Z\g<2>",
            members[core_path],
        )

    with zipfile.ZipFile(
        destination, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as archive:
        for name in sorted(members):
            info = zipfile.ZipInfo(name, date_time=_FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            archive.writestr(
                info,
                members[name],
                compress_type=zipfile.ZIP_DEFLATED,
                compresslevel=9,
            )


def _atomic_save_workbook(workbook: Workbook, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    raw_fd, raw_name = tempfile.mkstemp(
        prefix=f".{destination.name}.", suffix=".raw.tmp", dir=destination.parent
    )
    canonical_fd, canonical_name = tempfile.mkstemp(
        prefix=f".{destination.name}.", suffix=".tmp", dir=destination.parent
    )
    os.close(raw_fd)
    os.close(canonical_fd)
    raw_path = Path(raw_name)
    canonical_path = Path(canonical_name)
    try:
        # openpyxl otherwise refreshes core-document timestamps at save time.
        _set_fixed_properties(workbook)
        workbook.save(raw_path)
        _canonicalize_zip(raw_path, canonical_path)
        with canonical_path.open("rb") as handle:
            os.fsync(handle.fileno())
        os.replace(canonical_path, destination)
        _fsync_directory(destination.parent)
    finally:
        raw_path.unlink(missing_ok=True)
        canonical_path.unlink(missing_ok=True)


def _atomic_write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    temporary = Path(temporary_name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2, sort_keys=True)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
        _fsync_directory(path.parent)
    finally:
        temporary.unlink(missing_ok=True)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def materialize_workbooks(run_dir: str | Path, *, allow_partial: bool = False) -> dict[str, object]:
    """Materialize legacy feedback workbooks from validated stored responses.

    Every manifest row and response is preflighted before any workbook or workbook
    manifest is written. With ``allow_partial=False``, a missing or invalid record
    therefore leaves all existing materialized artifacts untouched.
    """

    root = Path(run_dir).resolve()
    requests = read_manifest(root / "manifest.csv")
    if not requests:
        raise WorkbookMaterializationError("Feedback manifest contains no requests")

    run_ids = {request.run_id for request in requests}
    model_ids = {request.model_id for request in requests}
    if len(run_ids) != 1 or len(model_ids) != 1:
        raise WorkbookMaterializationError(
            "Feedback manifest must contain exactly one run_id and one model_id"
        )

    grouped: dict[str, list[FeedbackRequest]] = defaultdict(list)
    for request in requests:
        validate_request_destinations(request)
        grouped[request.workbook_path].append(request)

    records: dict[str, Mapping[str, object]] = {}
    failures: dict[str, tuple[str, str]] = {}
    tables: dict[str, tuple[tuple[str, str, str, str], list[tuple[str, str, str, str]]]] = {}

    # Full preflight: paths, response identity/schema/hashes, payloads, and workbook shape.
    for workbook_path in sorted(grouped):
        workbook_requests = sorted(
            grouped[workbook_path], key=lambda item: (item.category_order, item.request_order)
        )
        _validate_sheet_contract(workbook_requests)
        _safe_run_path(root, workbook_path)
        for request in workbook_requests:
            response_path = _safe_run_path(root, request.response_path)
            if not response_path.is_file():
                failures[request.request_id] = ("missing", "response file does not exist")
                continue
            try:
                record = validate_response_record(request, response_path)
                if not isinstance(record, Mapping):
                    raise WorkbookMaterializationError(
                        f"{request.request_id}: validated response record must be an object"
                    )
                table = _table_for(request, record)
            except Exception:
                failures[request.request_id] = (
                    "invalid",
                    "response record failed validation",
                )
                continue
            records[request.request_id] = record
            tables[request.request_id] = table

    if failures and not allow_partial:
        failing_ids = ", ".join(
            request.request_id for request in requests if request.request_id in failures
        )
        raise WorkbookMaterializationError(
            f"Cannot materialize complete workbooks; missing or invalid requests: {failing_ids}"
        )

    workbook_entries: list[dict[str, object]] = []
    for workbook_path in sorted(
        grouped,
        key=lambda value: min(request.request_order for request in grouped[value]),
    ):
        workbook_requests = sorted(
            grouped[workbook_path], key=lambda item: (item.category_order, item.request_order)
        )
        missing_requests = [
            request for request in workbook_requests if request.request_id in failures
        ]
        usable_requests = [
            request for request in workbook_requests if request.request_id not in failures
        ]
        partial = bool(missing_requests)
        declared_path = _safe_run_path(root, workbook_path)
        destination = _partial_path(declared_path) if partial else declared_path

        if allow_partial and not usable_requests:
            declared_path.unlink(missing_ok=True)
            _partial_path(declared_path).unlink(missing_ok=True)
            continue

        workbook = Workbook()
        workbook.remove(workbook.active)
        _set_fixed_properties(workbook)
        dimensions: dict[str, dict[str, int]] = {}
        source_request_ids: list[str] = []

        for request in workbook_requests:
            if request.request_id in failures:
                continue
            worksheet = workbook.create_sheet(request.sheet_name)
            headers, rows = tables[request.request_id]
            _append_table(worksheet, headers, rows)
            dimensions[request.sheet_name] = {
                "rows": 6,
                "columns": 4,
                "data_rows": 5,
            }
            source_request_ids.append(request.request_id)

        if partial:
            worksheet = workbook.create_sheet(INCOMPLETE_SHEET_NAME)
            worksheet.append(("Request ID", "Status", "Reason"))
            for request in missing_requests:
                status, reason = failures[request.request_id]
                worksheet.append((request.request_id, status, reason))
            dimensions[INCOMPLETE_SHEET_NAME] = {
                "rows": len(missing_requests) + 1,
                "columns": 3,
                "data_rows": len(missing_requests),
            }

        _atomic_save_workbook(workbook, destination)
        alternate = declared_path if partial else _partial_path(declared_path)
        alternate.unlink(missing_ok=True)
        relative_destination = destination.relative_to(root).as_posix()
        digest = _sha256(destination)
        workbook_entries.append(
            {
                "workbook_path": relative_destination,
                "file_sha256": digest,
                "sha256": digest,
                "sheet_names": workbook.sheetnames,
                "dimensions": dimensions,
                "source_request_ids": source_request_ids,
                "model_id": workbook_requests[0].model_id,
                "run_id": workbook_requests[0].run_id,
                "partial": partial,
                "missing_or_invalid_request_ids": [
                    request.request_id for request in missing_requests
                ],
            }
        )

    manifest: dict[str, object] = {
        "schema_version": WORKBOOK_MANIFEST_SCHEMA_VERSION,
        "run_id": next(iter(run_ids)),
        "model_id": next(iter(model_ids)),
        "partial": bool(failures),
        "request_count": len(requests),
        "validated_request_count": len(records),
        "missing_or_invalid_request_ids": [
            request.request_id for request in requests if request.request_id in failures
        ],
        "workbooks": workbook_entries,
    }
    _atomic_write_json(root / "workbook_manifest.json", manifest)
    return manifest


__all__ = [
    "INCOMPLETE_SHEET_NAME",
    "WORKBOOK_MANIFEST_SCHEMA_VERSION",
    "WorkbookMaterializationError",
    "expected_workbook_path",
    "materialize_workbooks",
    "validate_request_destinations",
]
