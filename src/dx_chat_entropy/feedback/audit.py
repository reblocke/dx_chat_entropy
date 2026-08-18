from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import tempfile
import zipfile
from collections import Counter, defaultdict
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .manifest import LEDGER_COLUMNS, read_manifest
from .models import FeedbackRequest, canonical_json
from .storage import VALID_FAILURE_CLASSES, validate_attempt_record, validate_response_record
from .workbooks import (
    INCOMPLETE_SHEET_NAME,
    WORKBOOK_MANIFEST_SCHEMA_VERSION,
    validate_request_destinations,
)

QUALITY_SUMMARY_SCHEMA_VERSION = "feedback-quality-summary-v1"
INVALID_REQUEST_COLUMNS = (
    "request_order",
    "request_id",
    "check",
    "reason",
    "path",
)
LEDGER_STATUSES = {
    "pending",
    "running",
    "success",
    "invalid",
    "transient_failure",
    "permanent_failure",
    "skipped_existing",
}
_SUCCESS_LEDGER_STATUSES = {"success", "skipped_existing"}
_SECRET_PATTERNS = (
    ("openai_key", re.compile(r"sk-(?:proj-)?[A-Za-z0-9_-]{16,}")),
    ("github_token", re.compile(r"ghp_[A-Za-z0-9]{20,}")),
    ("slack_token", re.compile(r"xox[baprs]-[A-Za-z0-9-]{10,}")),
    ("authorization_header", re.compile(r"(?i)authorization\s*[:=]\s*bearer\s+\S+")),
    (
        "api_key_assignment",
        re.compile(r"(?i)api[_ -]?key\s*[:=]\s*[\"']?[A-Za-z0-9_-]{16,}"),
    ),
)


def _fsync_directory(directory: Path) -> None:
    try:
        descriptor = os.open(directory, os.O_RDONLY)
    except OSError:  # pragma: no cover - unsupported filesystems/platforms.
        return
    try:
        os.fsync(descriptor)
    except OSError:  # pragma: no cover - unsupported filesystems/platforms.
        pass
    finally:
        os.close(descriptor)


@dataclass(frozen=True)
class AuditFinding:
    request_order: int | None
    request_id: str
    check: str
    reason: str
    path: str = ""

    def __post_init__(self) -> None:
        # Audit findings are persisted and may also be rendered by callers. Keep
        # every string field behind the same final redaction boundary even when
        # a future check accidentally passes through untrusted artifact text.
        for field in ("request_id", "check", "reason", "path"):
            object.__setattr__(self, field, _safe_reason(getattr(self, field)))

    def row(self) -> dict[str, object]:
        return {
            "request_order": "" if self.request_order is None else self.request_order,
            "request_id": self.request_id,
            "check": self.check,
            "reason": self.reason,
            "path": self.path,
        }


@dataclass(frozen=True)
class AuditResult:
    passes: bool
    findings: tuple[AuditFinding, ...]
    summary: dict[str, object]

    def to_dict(self) -> dict[str, object]:
        return dict(self.summary)


def _safe_reason(value: object) -> str:
    text = str(value).replace("\r", " ").replace("\n", " ")
    for _name, pattern in _SECRET_PATTERNS:
        text = pattern.sub("[REDACTED]", text)
    return text[:1000]


def _safe_run_path(run_dir: Path, relative_path: str) -> Path:
    path = Path(relative_path)
    if path.is_absolute() or ".." in path.parts:
        raise ValueError(f"Run artifact path must be safe and relative: {relative_path!r}")
    resolved = (run_dir / path).resolve()
    resolved.relative_to(run_dir.resolve())
    return resolved


def _partial_path(path: Path) -> Path:
    return path.with_name(f"{path.stem}_PARTIAL{path.suffix}")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _sha256_json(value: object) -> str:
    return hashlib.sha256(canonical_json(value).encode("utf-8")).hexdigest()


def _inventory_category_mode(requests: Sequence[FeedbackRequest]) -> str | None:
    category_keys = {request.category_key for request in requests}
    if category_keys == {"subjective-and-historical"}:
        return "only_overall"
    if len(category_keys) > 1:
        return "all"
    return None


def _inventory_fingerprint(requests: Sequence[FeedbackRequest], category_mode: str) -> str | None:
    config_versions = {request.config_schema_version for request in requests}
    prompt_versions = {request.prompt_schema_version for request in requests}
    response_versions = {request.response_schema_version for request in requests}
    if not all(
        len(versions) == 1 for versions in (config_versions, prompt_versions, response_versions)
    ):
        return None
    payload = {
        "config_schema_version": next(iter(config_versions)),
        "prompt_schema_version": next(iter(prompt_versions)),
        "response_schema_version": next(iter(response_versions)),
        "category_mode": category_mode,
        "ordered_request_ids": [request.request_id for request in requests],
    }
    return _sha256_json(payload)


def _atomic_write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    temporary = Path(name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2, sort_keys=True)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
        _fsync_directory(path.parent)
    finally:
        temporary.unlink(missing_ok=True)


def _atomic_write_csv(path: Path, rows: Sequence[Mapping[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    temporary = Path(name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=INVALID_REQUEST_COLUMNS)
            writer.writeheader()
            writer.writerows(rows)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
        _fsync_directory(path.parent)
    finally:
        temporary.unlink(missing_ok=True)


def _response_table(
    request: FeedbackRequest, record: Mapping[str, object]
) -> tuple[list[str], list[list[str]]]:
    payload = record.get("parsed_payload")
    if not isinstance(payload, Mapping):
        raise ValueError("validated response record lacks parsed_payload")
    if request.surface.startswith("overall/"):
        first_key = "for_diagnosis_strongest_evidence"
        second_key = "against_diagnosis_strongest_evidence"
        diagnosis = request.target_diagnosis
        headers = [
            f"Supports {diagnosis}",
            f"Rationale for {diagnosis}",
            f"Against {diagnosis}",
            f"Rationale Against {diagnosis}",
        ]
    elif request.surface.startswith("differential/"):
        first_key = "diagnosisA_strongest_evidence"
        second_key = "diagnosisB_strongest_evidence"
        if not request.comparison_diagnosis:
            raise ValueError("differential request lacks comparison_diagnosis")
        headers = [
            f"Supports {request.correct_diagnosis}",
            f"Rationale {request.correct_diagnosis}",
            f"Supports {request.comparison_diagnosis}",
            f"Rationale {request.comparison_diagnosis}",
        ]
    else:
        raise ValueError(f"unsupported surface {request.surface!r}")

    first = payload.get(first_key)
    second = payload.get(second_key)
    if not isinstance(first, list) or not isinstance(second, list):
        raise ValueError("validated response evidence lists are missing")
    if len(first) != 5 or len(second) != 5:
        raise ValueError("validated response evidence lists must contain 5 items")

    rows: list[list[str]] = []
    for index in range(5):
        if not isinstance(first[index], Mapping) or not isinstance(second[index], Mapping):
            raise ValueError("validated response evidence item is not an object")
        rows.append(
            [
                str(first[index].get("finding", "")),
                str(first[index].get("explanation", "")),
                str(second[index].get("finding", "")),
                str(second[index].get("explanation", "")),
            ]
        )
    return headers, rows


def _scan_manifest_duplicates(path: Path) -> list[AuditFinding]:
    findings: list[AuditFinding] = []
    try:
        with path.open(encoding="utf-8", newline="") as handle:
            rows = list(csv.DictReader(handle))
    except Exception as exc:
        return [AuditFinding(None, "", "manifest_read", _safe_reason(exc), "manifest.csv")]

    for field, check in (("request_id", "duplicate_request"), ("request_order", "duplicate_order")):
        counts = Counter(row.get(field, "") for row in rows)
        for value in sorted(item for item, count in counts.items() if item and count > 1):
            findings.append(
                AuditFinding(
                    None,
                    value if field == "request_id" else "",
                    check,
                    f"duplicate {field}: {value}",
                    "manifest.csv",
                )
            )
    destinations = Counter(
        (row.get("workbook_path", ""), row.get("sheet_name", "")) for row in rows
    )
    for (workbook_path, sheet_name), count in sorted(destinations.items()):
        if workbook_path and sheet_name and count > 1:
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "duplicate_destination",
                    f"duplicate workbook/sheet destination: {workbook_path}#{sheet_name}",
                    "manifest.csv",
                )
            )
    return findings


def _scan_response_duplicates(
    run_dir: Path, requests: Sequence[FeedbackRequest]
) -> list[AuditFinding]:
    findings: list[AuditFinding] = []
    expected_ids = {request.request_id for request in requests}
    order_by_id = {request.request_id: request.request_order for request in requests}
    found: dict[str, list[Path]] = defaultdict(list)
    responses_dir = run_dir / "responses"
    if not responses_dir.exists():
        return findings

    for path in sorted(item for item in responses_dir.rglob("*") if item.is_file()):
        relative = path.relative_to(run_dir).as_posix()
        if path.suffix != ".json":
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "unexpected_response_file",
                    "responses contains an unexpected non-JSON file",
                    relative,
                )
            )
            continue
        try:
            value = json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "response_json",
                    _safe_reason(exc),
                    relative,
                )
            )
            continue
        request_id = value.get("request_id") if isinstance(value, Mapping) else None
        if not isinstance(request_id, str) or not request_id:
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "response_identity",
                    "response has no request_id",
                    relative,
                )
            )
            continue
        found[request_id].append(path)
        if request_id not in expected_ids:
            findings.append(
                AuditFinding(
                    None,
                    request_id,
                    "unexpected_response",
                    "response request_id is not in manifest",
                    relative,
                )
            )

    for request_id, paths in sorted(found.items()):
        if len(paths) > 1:
            findings.append(
                AuditFinding(
                    order_by_id.get(request_id),
                    request_id,
                    "duplicate_response",
                    f"found {len(paths)} response files for request_id",
                    ";".join(path.relative_to(run_dir).as_posix() for path in paths),
                )
            )
    return findings


def _audit_attempts(
    run_dir: Path,
    requests: Sequence[FeedbackRequest],
    records: Mapping[str, Mapping[str, object]],
    ledger_rows: Mapping[str, Mapping[str, str]],
) -> list[AuditFinding]:
    findings: list[AuditFinding] = []
    requests_by_id = {request.request_id: request for request in requests}
    attempts_dir = run_dir / "attempts"
    if not attempts_dir.is_dir():
        return [
            AuditFinding(None, "", "attempt_inventory", "attempts directory is missing", "attempts")
        ]

    found: dict[str, dict[int, Path]] = defaultdict(dict)
    normalized_records: dict[str, dict[int, dict[str, object]]] = defaultdict(dict)
    for path in sorted(item for item in attempts_dir.rglob("*") if item.is_file()):
        relative = path.relative_to(run_dir).as_posix()
        try:
            parts = path.relative_to(attempts_dir).parts
            if len(parts) != 2 or path.suffix != ".json":
                raise ValueError("attempt record path differs from the required contract")
            request_id = parts[0]
            attempt_text = path.stem
            if not re.fullmatch(r"[0-9]{4}", attempt_text):
                raise ValueError("attempt filename must be a zero-padded integer")
            attempt = int(attempt_text)
            request = requests_by_id.get(request_id)
            if request is None:
                raise ValueError("attempt request_id is not in the manifest")
            if attempt in found[request_id]:
                raise ValueError("duplicate attempt number")
            found[request_id][attempt] = path
            normalized = validate_attempt_record(
                path,
                expected_request_id=request_id,
                expected_attempt=attempt,
            )
            normalized_records[request_id][attempt] = normalized
        except json.JSONDecodeError as exc:  # pragma: no cover - storage wraps JSON errors.
            findings.append(AuditFinding(None, "", "attempt_json", _safe_reason(exc), relative))
        except Exception as exc:
            findings.append(AuditFinding(None, "", "attempt_schema", _safe_reason(exc), relative))

    for request in requests:
        row = ledger_rows.get(request.request_id)
        if row is None:
            continue
        try:
            attempt_count = int(row.get("attempt_count", ""))
        except (TypeError, ValueError):
            continue
        actual_attempts = sorted(found.get(request.request_id, {}))
        successful_attempts = sorted(
            attempt
            for attempt, record in normalized_records.get(request.request_id, {}).items()
            if record["status"] == "success"
        )
        if len(successful_attempts) > 1:
            second_success = successful_attempts[1]
            findings.append(
                AuditFinding(
                    request.request_order,
                    request.request_id,
                    "attempt_history",
                    "attempt history contains more than one successful attempt",
                    found[request.request_id][second_success].relative_to(run_dir).as_posix(),
                )
            )
        if successful_attempts:
            first_success = successful_attempts[0]
            for attempt in (item for item in actual_attempts if item > first_success):
                findings.append(
                    AuditFinding(
                        request.request_order,
                        request.request_id,
                        "attempt_history",
                        "attempt history continues after a successful attempt",
                        found[request.request_id][attempt].relative_to(run_dir).as_posix(),
                    )
                )
        expected_attempts = list(range(1, attempt_count + 1))
        if actual_attempts != expected_attempts:
            unexpected_attempts = sorted(set(actual_attempts) - set(expected_attempts))
            inventory_path = (
                found[request.request_id][unexpected_attempts[0]].relative_to(run_dir).as_posix()
                if unexpected_attempts
                else f"attempts/{request.request_id}"
            )
            findings.append(
                AuditFinding(
                    request.request_order,
                    request.request_id,
                    "attempt_inventory",
                    "attempt files do not match the ledger attempt_count",
                    inventory_path,
                )
            )
        if attempt_count > 0 and attempt_count in normalized_records.get(request.request_id, {}):
            final_status = str(normalized_records[request.request_id][attempt_count]["status"])
            ledger_status = str(row.get("status", ""))
            expected_final_status = (
                "success" if ledger_status in _SUCCESS_LEDGER_STATUSES else ledger_status
            )
            if final_status != expected_final_status:
                findings.append(
                    AuditFinding(
                        request.request_order,
                        request.request_id,
                        "attempt_status",
                        "final attempt status differs from the ledger state",
                        found[request.request_id][attempt_count].relative_to(run_dir).as_posix(),
                    )
                )
        record = records.get(request.request_id)
        if record is not None and record.get("attempt_count") != attempt_count:
            findings.append(
                AuditFinding(
                    request.request_order,
                    request.request_id,
                    "attempt_inventory",
                    "response attempt_count differs from the ledger",
                    request.response_path,
                )
            )
    return findings


def _read_ledger_for_attempt_audit(run_dir: Path) -> dict[str, dict[str, str]]:
    path = run_dir / "run_ledger.csv"
    if not path.is_file():
        return {}
    try:
        with path.open(encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle)
            if tuple(reader.fieldnames or ()) != LEDGER_COLUMNS:
                return {}
            return {row.get("request_id", ""): row for row in reader}
    except Exception:
        return {}


def _audit_ledger(
    run_dir: Path,
    requests: Sequence[FeedbackRequest],
    valid_ids: set[str],
) -> tuple[list[AuditFinding], Counter[str]]:
    path = run_dir / "run_ledger.csv"
    if not path.is_file():
        return [
            AuditFinding(None, "", "ledger_missing", "run_ledger.csv is missing", "run_ledger.csv")
        ], Counter()

    findings: list[AuditFinding] = []
    statuses: Counter[str] = Counter()
    request_by_id = {request.request_id: request for request in requests}
    try:
        with path.open(encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle)
            if tuple(reader.fieldnames or ()) != LEDGER_COLUMNS:
                findings.append(
                    AuditFinding(
                        None,
                        "",
                        "ledger_schema",
                        "ledger columns differ from the required contract",
                        "run_ledger.csv",
                    )
                )
            rows = list(reader)
    except Exception as exc:
        return [
            AuditFinding(None, "", "ledger_read", _safe_reason(exc), "run_ledger.csv")
        ], statuses

    by_id: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        request_id = row.get("request_id", "")
        by_id[request_id].append(row)
    if [row.get("request_id", "") for row in rows] != [request.request_id for request in requests]:
        findings.append(
            AuditFinding(
                None,
                "",
                "ledger_order",
                "ledger row order differs from the immutable manifest",
                "run_ledger.csv",
            )
        )
    for request_id, request in request_by_id.items():
        matching = by_id.get(request_id, [])
        if len(matching) != 1:
            findings.append(
                AuditFinding(
                    request.request_order,
                    request_id,
                    "ledger_cardinality",
                    f"expected one ledger row, found {len(matching)}",
                    "run_ledger.csv",
                )
            )
            continue
        row = matching[0]
        status = row.get("status", "")
        if status in LEDGER_STATUSES:
            statuses[status] += 1
        else:
            statuses["invalid_status"] += 1
        if row.get("request_order") != str(request.request_order):
            findings.append(
                AuditFinding(
                    request.request_order,
                    request_id,
                    "ledger_order",
                    "ledger request_order differs from manifest",
                    "run_ledger.csv",
                )
            )
        try:
            attempt_count = int(row.get("attempt_count", ""))
            if attempt_count < 0:
                raise ValueError
        except (TypeError, ValueError):
            findings.append(
                AuditFinding(
                    request.request_order,
                    request_id,
                    "ledger_attempt_count",
                    "ledger attempt_count must be a non-negative integer",
                    "run_ledger.csv",
                )
            )
        if status not in LEDGER_STATUSES:
            findings.append(
                AuditFinding(
                    request.request_order,
                    request_id,
                    "ledger_status",
                    "ledger status is not allowlisted",
                    "run_ledger.csv",
                )
            )
        failure_class = row.get("failure_class", "")
        if failure_class not in VALID_FAILURE_CLASSES:
            findings.append(
                AuditFinding(
                    request.request_order,
                    request_id,
                    "ledger_failure_class",
                    "ledger failure_class is not allowlisted",
                    "run_ledger.csv",
                )
            )
        http_status = row.get("http_status", "")
        if http_status != "" and (
            not isinstance(http_status, str)
            or not http_status.isascii()
            or not http_status.isdecimal()
            or not 100 <= int(http_status) <= 599
        ):
            findings.append(
                AuditFinding(
                    request.request_order,
                    request_id,
                    "ledger_http_status",
                    "ledger http_status must be blank or an integer from 100 through 599",
                    "run_ledger.csv",
                )
            )
        updated_at = row.get("updated_at", "")
        try:
            if not updated_at:
                raise ValueError
            if isinstance(updated_at, str):
                parsed_updated_at = datetime.fromisoformat(updated_at.replace("Z", "+00:00"))
            else:
                raise ValueError
            if parsed_updated_at.tzinfo is None or parsed_updated_at.utcoffset() is None:
                raise ValueError
        except (TypeError, ValueError):
            findings.append(
                AuditFinding(
                    request.request_order,
                    request_id,
                    "ledger_updated_at",
                    "ledger updated_at must be an ISO 8601 timestamp with UTC offset",
                    "run_ledger.csv",
                )
            )
        if status == "running":
            findings.append(
                AuditFinding(
                    request.request_order,
                    request_id,
                    "stale_running",
                    "ledger still marks request running",
                    "run_ledger.csv",
                )
            )
        if (status in _SUCCESS_LEDGER_STATUSES) != (request_id in valid_ids):
            findings.append(
                AuditFinding(
                    request.request_order,
                    request_id,
                    "ledger_response_mismatch",
                    "ledger status disagrees with validated response state",
                    "run_ledger.csv",
                )
            )
        if row.get("response_path", "") != request.response_path:
            findings.append(
                AuditFinding(
                    request.request_order,
                    request_id,
                    "ledger_response_path",
                    "ledger response_path differs from manifest",
                    "run_ledger.csv",
                )
            )

    for request_id in sorted(set(by_id) - set(request_by_id)):
        findings.append(
            AuditFinding(
                None,
                request_id,
                "unexpected_ledger_row",
                "ledger request_id is not in manifest",
                "run_ledger.csv",
            )
        )
    return findings, statuses


def _audit_manifest_summary(
    run_dir: Path, requests: Sequence[FeedbackRequest]
) -> list[AuditFinding]:
    path = run_dir / "manifest_summary.json"
    if not path.is_file():
        return [
            AuditFinding(
                None,
                "",
                "manifest_summary_missing",
                "manifest_summary.json is missing",
                "manifest_summary.json",
            )
        ]
    try:
        summary = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        return [
            AuditFinding(
                None,
                "",
                "manifest_summary_json",
                _safe_reason(exc),
                "manifest_summary.json",
            )
        ]
    if not isinstance(summary, Mapping):
        return [
            AuditFinding(
                None,
                "",
                "manifest_summary_schema",
                "manifest summary must be an object",
                "manifest_summary.json",
            )
        ]

    run_ids = {request.run_id for request in requests}
    category_mode = _inventory_category_mode(requests)
    fingerprint = (
        _inventory_fingerprint(requests, category_mode) if category_mode is not None else None
    )
    expected: dict[str, object] = {
        "schema_version": "feedback-manifest-summary-v1",
        "run_id": next(iter(run_ids)) if len(run_ids) == 1 else None,
        "run_fingerprint": fingerprint,
        "category_mode": category_mode,
        "total_requests": len(requests),
        "counts_by_surface": dict(Counter(request.surface for request in requests)),
        "counts_by_category": dict(Counter(request.category_key for request in requests)),
    }
    findings: list[AuditFinding] = []
    if set(summary) != set(expected):
        findings.append(
            AuditFinding(
                None,
                "",
                "manifest_summary_schema",
                "manifest summary fields differ from the required contract",
                "manifest_summary.json",
            )
        )
    for field, expected_value in expected.items():
        if summary.get(field) != expected_value:
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "manifest_summary_identity",
                    f"manifest summary {field} differs from the manifest contract",
                    "manifest_summary.json",
                )
            )
    return findings


def _audit_run_metadata(run_dir: Path, requests: Sequence[FeedbackRequest]) -> list[AuditFinding]:
    path = run_dir / "run_metadata.json"
    if not path.is_file():
        return [
            AuditFinding(
                None,
                "",
                "run_metadata_missing",
                "run_metadata.json is missing",
                "run_metadata.json",
            )
        ]
    try:
        metadata = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        return [
            AuditFinding(
                None,
                "",
                "run_metadata_json",
                _safe_reason(exc),
                "run_metadata.json",
            )
        ]
    if not isinstance(metadata, Mapping):
        return [
            AuditFinding(
                None,
                "",
                "run_metadata_schema",
                "run metadata must be an object",
                "run_metadata.json",
            )
        ]

    findings: list[AuditFinding] = []
    request_ids = [request.request_id for request in requests]
    run_ids = {request.run_id for request in requests}
    model_ids = {request.model_id for request in requests}
    model_profiles = {request.model_profile for request in requests}
    model_settings = {canonical_json(request.model_settings) for request in requests}
    config_versions = {request.config_schema_version for request in requests}
    prompt_versions = {request.prompt_schema_version for request in requests}
    response_versions = {request.response_schema_version for request in requests}
    category_mode = _inventory_category_mode(requests)
    fingerprint = (
        _inventory_fingerprint(requests, category_mode) if category_mode is not None else None
    )
    manifest_path = run_dir / "manifest.csv"
    expected: dict[str, object] = {
        "schema_version": "feedback-run-metadata-v1",
        "run_id": next(iter(run_ids)) if len(run_ids) == 1 else None,
        "run_fingerprint": fingerprint,
        "config_schema_version": (
            next(iter(config_versions)) if len(config_versions) == 1 else None
        ),
        "prompt_schema_version": (
            next(iter(prompt_versions)) if len(prompt_versions) == 1 else None
        ),
        "response_schema_version": (
            next(iter(response_versions)) if len(response_versions) == 1 else None
        ),
        "category_mode": category_mode,
        "model_profile": next(iter(model_profiles)) if len(model_profiles) == 1 else None,
        "model_id": next(iter(model_ids)) if len(model_ids) == 1 else None,
        "model_settings": requests[0].model_settings if len(model_settings) == 1 else None,
        "request_count": len(requests),
        "ordered_request_ids": request_ids,
        "manifest_sha256": _sha256(manifest_path) if manifest_path.is_file() else None,
        "manifest_path": "manifest.csv",
        "summary_path": "manifest_summary.json",
        "ledger_path": "run_ledger.csv",
    }
    if set(metadata) != set(expected):
        findings.append(
            AuditFinding(
                None,
                "",
                "run_metadata_schema",
                "run metadata fields differ from the required contract",
                "run_metadata.json",
            )
        )
    for field, expected_value in expected.items():
        if metadata.get(field) != expected_value:
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "run_metadata_identity",
                    f"run metadata {field} differs from manifest contract",
                    "run_metadata.json",
                )
            )

    if fingerprint is not None and len(run_ids) == 1:
        base_run_id = f"fb_{fingerprint[:24]}"
        run_id = next(iter(run_ids))
        if re.fullmatch(rf"{re.escape(base_run_id)}(?:_r[0-9]{{3}})?", run_id) is None:
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "run_identity",
                    "run_id does not match the manifest fingerprint",
                    "manifest.csv",
                )
            )
        if run_dir.name != run_id:
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "run_identity",
                    "run directory name differs from manifest run_id",
                    ".",
                )
            )
    return findings


def _audit_zip_determinism(path: Path) -> str | None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
    except Exception as exc:
        return f"invalid XLSX ZIP: {_safe_reason(exc)}"
    names = [info.filename for info in infos]
    if names != sorted(names):
        return "XLSX ZIP members are not in deterministic sorted order"
    if any(info.date_time != (1980, 1, 1, 0, 0, 0) for info in infos):
        return "XLSX ZIP members do not use the fixed deterministic timestamp"
    return None


def _audit_workbooks(
    run_dir: Path,
    requests: Sequence[FeedbackRequest],
    records: Mapping[str, Mapping[str, object]],
) -> tuple[list[AuditFinding], int]:
    findings: list[AuditFinding] = []
    path = run_dir / "workbook_manifest.json"
    if not path.is_file():
        return [
            AuditFinding(
                None,
                "",
                "workbook_manifest_missing",
                "workbook_manifest.json is missing",
                "workbook_manifest.json",
            )
        ], 0
    try:
        manifest = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        return [
            AuditFinding(
                None, "", "workbook_manifest_json", _safe_reason(exc), "workbook_manifest.json"
            )
        ], 0
    if not isinstance(manifest, Mapping):
        return [
            AuditFinding(
                None,
                "",
                "workbook_manifest_schema",
                "workbook manifest must be an object",
                "workbook_manifest.json",
            )
        ], 0

    run_ids = {request.run_id for request in requests}
    model_ids = {request.model_id for request in requests}
    if manifest.get("schema_version") != WORKBOOK_MANIFEST_SCHEMA_VERSION:
        findings.append(
            AuditFinding(
                None,
                "",
                "workbook_manifest_schema",
                "unexpected workbook manifest schema_version",
                "workbook_manifest.json",
            )
        )
    if len(run_ids) == 1 and manifest.get("run_id") != next(iter(run_ids)):
        findings.append(
            AuditFinding(
                None,
                "",
                "workbook_manifest_identity",
                "workbook manifest run_id differs from request manifest",
                "workbook_manifest.json",
            )
        )
    if len(model_ids) == 1 and manifest.get("model_id") != next(iter(model_ids)):
        findings.append(
            AuditFinding(
                None,
                "",
                "workbook_manifest_identity",
                "workbook manifest model_id differs from request manifest",
                "workbook_manifest.json",
            )
        )

    invalid_request_ids = [
        request.request_id for request in requests if request.request_id not in records
    ]
    top_level_contract = {
        "partial": bool(invalid_request_ids),
        "request_count": len(requests),
        "validated_request_count": len(records),
        "missing_or_invalid_request_ids": invalid_request_ids,
    }
    for field, expected_value in top_level_contract.items():
        if manifest.get(field) != expected_value:
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "workbook_manifest_contract",
                    f"workbook manifest {field} differs from validated run state",
                    "workbook_manifest.json",
                )
            )

    raw_entries = manifest.get("workbooks")
    if not isinstance(raw_entries, list):
        findings.append(
            AuditFinding(
                None,
                "",
                "workbook_manifest_schema",
                "workbooks must be an array",
                "workbook_manifest.json",
            )
        )
        raw_entries = []
    entries = {
        entry.get("workbook_path"): entry
        for entry in raw_entries
        if isinstance(entry, Mapping) and isinstance(entry.get("workbook_path"), str)
    }
    if len(entries) != len(raw_entries):
        findings.append(
            AuditFinding(
                None,
                "",
                "duplicate_workbook",
                "workbook manifest has duplicate or malformed entries",
                "workbook_manifest.json",
            )
        )

    grouped: dict[str, list[FeedbackRequest]] = defaultdict(list)
    for request in requests:
        grouped[request.workbook_path].append(request)

    audited = 0
    expected_paths: set[str] = set()
    for declared, group in sorted(
        grouped.items(), key=lambda item: min(request.request_order for request in item[1])
    ):
        group = sorted(group, key=lambda request: (request.category_order, request.request_order))
        invalid = [request for request in group if request.request_id not in records]
        expected_valid = [request for request in group if request.request_id in records]
        try:
            declared_path = _safe_run_path(run_dir, declared)
        except Exception as exc:
            findings.append(AuditFinding(None, "", "workbook_path", _safe_reason(exc), declared))
            continue
        if not expected_valid:
            continue
        actual_path = _partial_path(declared_path) if invalid else declared_path
        relative = actual_path.relative_to(run_dir).as_posix()
        expected_paths.add(relative)
        other_path = declared_path if invalid else _partial_path(declared_path)
        if other_path.exists():
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "duplicate_workbook",
                    "stale conflicting complete/partial workbook exists",
                    other_path.relative_to(run_dir).as_posix(),
                )
            )
        if not actual_path.is_file():
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "workbook_missing",
                    "expected materialized workbook is missing",
                    relative,
                )
            )
            continue
        audited += 1
        entry = entries.get(relative)
        if not isinstance(entry, Mapping):
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "workbook_manifest_entry",
                    "workbook is absent from workbook manifest",
                    relative,
                )
            )
        else:
            actual_hash = _sha256(actual_path)
            declared_hashes = [
                entry.get(field) for field in ("file_sha256", "sha256") if field in entry
            ]
            if not declared_hashes or any(value != actual_hash for value in declared_hashes):
                findings.append(
                    AuditFinding(
                        None,
                        "",
                        "workbook_hash",
                        "workbook hash differs from workbook manifest",
                        relative,
                    )
                )

        deterministic_error = _audit_zip_determinism(actual_path)
        if deterministic_error:
            findings.append(
                AuditFinding(None, "", "workbook_determinism", deterministic_error, relative)
            )

        expected_sheets = [request.sheet_name for request in expected_valid]
        expected_dimensions: dict[str, dict[str, int]] = {
            request.sheet_name: {"rows": 6, "columns": 4, "data_rows": 5}
            for request in expected_valid
        }
        if invalid:
            expected_sheets.append(INCOMPLETE_SHEET_NAME)
            expected_dimensions[INCOMPLETE_SHEET_NAME] = {
                "rows": len(invalid) + 1,
                "columns": 3,
                "data_rows": len(invalid),
            }
        if isinstance(entry, Mapping):
            entry_contract: dict[str, object] = {
                "sheet_names": expected_sheets,
                "dimensions": expected_dimensions,
                "source_request_ids": [request.request_id for request in expected_valid],
                "model_id": group[0].model_id,
                "run_id": group[0].run_id,
                "partial": bool(invalid),
                "missing_or_invalid_request_ids": [request.request_id for request in invalid],
            }
            for field, expected_value in entry_contract.items():
                if entry.get(field) != expected_value:
                    findings.append(
                        AuditFinding(
                            None,
                            "",
                            "workbook_manifest_contract",
                            f"workbook entry {field} differs from the workbook contract",
                            relative,
                        )
                    )
        try:
            workbook = load_workbook(actual_path, read_only=True, data_only=False)
            if workbook.sheetnames != expected_sheets:
                findings.append(
                    AuditFinding(
                        None,
                        "",
                        "workbook_sheets",
                        f"expected sheets {expected_sheets!r}, found {workbook.sheetnames!r}",
                        relative,
                    )
                )
            for request in expected_valid:
                if request.sheet_name not in workbook.sheetnames:
                    continue
                worksheet = workbook[request.sheet_name]
                expected_headers, expected_rows = _response_table(
                    request, records[request.request_id]
                )
                actual = [list(row) for row in worksheet.iter_rows(values_only=True)]
                if worksheet.max_row != 6 or worksheet.max_column != 4:
                    findings.append(
                        AuditFinding(
                            request.request_order,
                            request.request_id,
                            "workbook_dimensions",
                            "feedback sheet must have one header row, 5 data rows, and 4 columns",
                            relative,
                        )
                    )
                if actual != [expected_headers, *expected_rows]:
                    findings.append(
                        AuditFinding(
                            request.request_order,
                            request.request_id,
                            "workbook_values",
                            "feedback sheet values differ from validated response",
                            relative,
                        )
                    )
            if invalid and workbook.sheetnames and workbook.sheetnames[-1] == INCOMPLETE_SHEET_NAME:
                worksheet = workbook[INCOMPLETE_SHEET_NAME]
                listed_ids = [row[0] for row in list(worksheet.iter_rows(values_only=True))[1:]]
                expected_ids = [request.request_id for request in invalid]
                if listed_ids != expected_ids:
                    findings.append(
                        AuditFinding(
                            None,
                            "",
                            "partial_label",
                            "INCOMPLETE sheet does not list missing/invalid request IDs "
                            "in manifest order",
                            relative,
                        )
                    )
            workbook.close()
        except Exception as exc:
            findings.append(AuditFinding(None, "", "workbook_read", _safe_reason(exc), relative))

    for extra in sorted(set(entries) - expected_paths):
        findings.append(
            AuditFinding(
                None,
                "",
                "unexpected_workbook",
                "workbook manifest entry is not expected by request manifest",
                str(extra),
            )
        )
    physical_paths = {
        workbook.relative_to(run_dir).as_posix()
        for workbook in (run_dir / "workbooks").rglob("*.xlsx")
    }
    for extra in sorted(physical_paths - expected_paths):
        findings.append(
            AuditFinding(
                None,
                "",
                "unexpected_workbook",
                "physical workbook is not expected by validated source responses",
                extra,
            )
        )
    return findings, audited


def _scan_temporary_files(run_dir: Path) -> list[AuditFinding]:
    findings: list[AuditFinding] = []
    for path in sorted(item for item in run_dir.rglob("*") if item.is_file()):
        name = path.name.lower()
        if name.endswith((".tmp", ".partial")) or ".raw.tmp" in name:
            findings.append(
                AuditFinding(
                    None,
                    "",
                    "temporary_file",
                    "stale temporary artifact",
                    path.relative_to(run_dir).as_posix(),
                )
            )
    return findings


def _scan_secrets(run_dir: Path) -> list[AuditFinding]:
    findings: list[AuditFinding] = []
    excluded = {"invalid_requests.csv", "quality_summary.json"}
    for path in sorted(
        item for item in run_dir.rglob("*") if item.is_file() and item.name not in excluded
    ):
        try:
            if path.suffix.lower() == ".xlsx":
                with zipfile.ZipFile(path) as archive:
                    text = "\n".join(
                        archive.read(name).decode("utf-8", errors="ignore")
                        for name in archive.namelist()
                    )
            else:
                text = path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            continue
        for name, pattern in _SECRET_PATTERNS:
            if pattern.search(text):
                findings.append(
                    AuditFinding(
                        None,
                        "",
                        "secret_pattern",
                        f"artifact contains {name}",
                        path.relative_to(run_dir).as_posix(),
                    )
                )
    return findings


def _finding_sort_key(finding: AuditFinding) -> tuple[int, str, str, str, str]:
    order = finding.request_order if finding.request_order is not None else 2**63 - 1
    return (order, finding.request_id, finding.check, finding.path, finding.reason)


def audit_feedback_run(run_dir: str | Path) -> AuditResult:
    """Audit a feedback run and write deterministic repair and quality reports."""

    root = Path(run_dir).resolve()
    findings: list[AuditFinding] = []
    manifest_path = root / "manifest.csv"
    requests: list[FeedbackRequest] = []
    records: dict[str, Mapping[str, object]] = {}

    if not manifest_path.is_file():
        findings.append(
            AuditFinding(None, "", "manifest_missing", "manifest.csv is missing", "manifest.csv")
        )
    else:
        findings.extend(_scan_manifest_duplicates(manifest_path))
        try:
            requests = read_manifest(manifest_path)
        except Exception as exc:
            findings.append(
                AuditFinding(None, "", "manifest_validation", _safe_reason(exc), "manifest.csv")
            )

    run_ids = {request.run_id for request in requests}
    model_ids = {request.model_id for request in requests}
    if requests and len(run_ids) != 1:
        findings.append(
            AuditFinding(
                None, "", "run_identity", "manifest contains multiple run IDs", "manifest.csv"
            )
        )
    if requests and len(model_ids) != 1:
        findings.append(
            AuditFinding(
                None, "", "model_identity", "manifest contains multiple model IDs", "manifest.csv"
            )
        )

    if requests:
        findings.extend(_audit_manifest_summary(root, requests))
        findings.extend(_audit_run_metadata(root, requests))

    for request in requests:
        try:
            validate_request_destinations(request)
        except Exception as exc:
            findings.append(
                AuditFinding(
                    request.request_order,
                    request.request_id,
                    "artifact_contract",
                    _safe_reason(exc),
                    "manifest.csv",
                )
            )
        try:
            response_path = _safe_run_path(root, request.response_path)
        except Exception as exc:
            findings.append(
                AuditFinding(
                    request.request_order,
                    request.request_id,
                    "response_path",
                    _safe_reason(exc),
                    request.response_path,
                )
            )
            continue
        if not response_path.is_file():
            findings.append(
                AuditFinding(
                    request.request_order,
                    request.request_id,
                    "missing_response",
                    "response file does not exist",
                    request.response_path,
                )
            )
            continue
        try:
            record = validate_response_record(request, response_path)
            if not isinstance(record, Mapping):
                raise ValueError("validated response record must be an object")
            _response_table(request, record)
            records[request.request_id] = record
        except Exception as exc:
            findings.append(
                AuditFinding(
                    request.request_order,
                    request.request_id,
                    "invalid_response",
                    _safe_reason(exc),
                    request.response_path,
                )
            )

    findings.extend(_scan_response_duplicates(root, requests))
    ledger_rows = _read_ledger_for_attempt_audit(root)
    ledger_findings, ledger_statuses = (
        _audit_ledger(root, requests, set(records)) if requests else ([], Counter())
    )
    findings.extend(ledger_findings)
    if requests:
        findings.extend(_audit_attempts(root, requests, records, ledger_rows))
    workbook_findings, workbook_count = (
        _audit_workbooks(root, requests, records) if requests else ([], 0)
    )
    findings.extend(workbook_findings)
    findings.extend(_scan_temporary_files(root))
    findings.extend(_scan_secrets(root))
    findings = sorted(set(findings), key=_finding_sort_key)

    invalid_ids = {finding.request_id for finding in findings if finding.request_id}
    checks = Counter(finding.check for finding in findings)
    summary: dict[str, Any] = {
        "schema_version": QUALITY_SUMMARY_SCHEMA_VERSION,
        "passes": not findings,
        "run_id": next(iter(run_ids)) if len(run_ids) == 1 else None,
        "model_id": next(iter(model_ids)) if len(model_ids) == 1 else None,
        "manifest_request_count": len(requests),
        "validated_response_count": len(records),
        "invalid_request_count": len(invalid_ids),
        "finding_count": len(findings),
        "workbook_count": workbook_count,
        "ledger_status_counts": dict(sorted(ledger_statuses.items())),
        "finding_counts": dict(sorted(checks.items())),
        "manifest_sha256": _sha256(manifest_path) if manifest_path.is_file() else None,
    }
    _atomic_write_csv(root / "invalid_requests.csv", [finding.row() for finding in findings])
    _atomic_write_json(root / "quality_summary.json", summary)
    return AuditResult(passes=not findings, findings=tuple(findings), summary=summary)


__all__ = [
    "AuditFinding",
    "AuditResult",
    "INVALID_REQUEST_COLUMNS",
    "QUALITY_SUMMARY_SCHEMA_VERSION",
    "audit_feedback_run",
]
