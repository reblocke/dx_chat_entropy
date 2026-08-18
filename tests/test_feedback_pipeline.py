from __future__ import annotations

import csv
import hashlib
import json
import runpy
import socket
import time
import zipfile
from collections import Counter
from collections.abc import Callable, Mapping
from copy import deepcopy
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from dx_chat_entropy.feedback.audit import audit_feedback_run
from dx_chat_entropy.feedback.config import FeedbackConfig, load_feedback_config
from dx_chat_entropy.feedback.manifest import (
    build_feedback_requests,
    build_or_load_manifest,
    prompt_sha256,
    read_manifest,
)
from dx_chat_entropy.feedback.models import (
    MANIFEST_COLUMNS,
    DifferentialFeedbackResponse,
    FeedbackRequest,
    OverallFeedbackResponse,
    ProviderResult,
    canonical_json,
    validate_response_payload,
)
from dx_chat_entropy.feedback.runtime import (
    FakeProviderAdapter,
    OpenAIProviderAdapter,
    execute_feedback_run,
    select_requests,
)
from dx_chat_entropy.feedback.storage import (
    RESPONSE_RECORD_SCHEMA_VERSION,
    LedgerStore,
    atomic_write_bytes,
    create_recompute_run,
    validate_attempt_record,
    validate_response_record,
    write_attempt_record,
    write_response_record,
)
from dx_chat_entropy.feedback.workbooks import (
    INCOMPLETE_SHEET_NAME,
    WorkbookMaterializationError,
    materialize_workbooks,
)

REPO_ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = REPO_ROOT / "config" / "feedback_generation.yaml"
CONTRACT_PATH = REPO_ROOT / "tests" / "fixtures" / "feedback" / "notebook_contract.json"
SECRET_SENTINEL = "sk-proj-THIS_MUST_NEVER_REACH_AN_ARTIFACT_0123456789"


@pytest.fixture(autouse=True)
def block_feedback_test_network(monkeypatch: pytest.MonkeyPatch) -> None:
    """Make every feedback test fail immediately on an accidental network call."""

    def denied(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("feedback tests must not open sockets")

    monkeypatch.setattr(socket, "create_connection", denied)
    monkeypatch.setattr(socket.socket, "connect", denied)
    monkeypatch.setattr(socket.socket, "connect_ex", denied)


@pytest.fixture(scope="module")
def contract() -> dict[str, object]:
    value = json.loads(CONTRACT_PATH.read_text(encoding="utf-8"))
    assert isinstance(value, dict)
    return value


@pytest.fixture(scope="module")
def feedback_config() -> FeedbackConfig:
    return load_feedback_config(CONFIG_PATH)


def _ordered_messages_sha256(requests: tuple[FeedbackRequest, ...]) -> str:
    value = canonical_json([message for request in requests for message in request.messages])
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _raw_config() -> dict[str, object]:
    value = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    assert isinstance(value, dict)
    return value


def _write_raw_config(tmp_path: Path, value: Mapping[str, object]) -> Path:
    path = tmp_path / "feedback_generation.yaml"
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def _payload_for(request: FeedbackRequest, *, marker: str = "fake") -> dict[str, object]:
    items_a = [
        {
            "finding": f"{marker} A finding {index}",
            "explanation": f"{marker} A explanation {index}",
            "abbreviation_expansion": {},
        }
        for index in range(1, 6)
    ]
    items_b = [
        {
            "finding": f"{marker} B finding {index}",
            "explanation": f"{marker} B explanation {index}",
            "abbreviation_expansion": {},
        }
        for index in range(1, 6)
    ]
    if request.expected_response_type == "overall":
        return {
            "for_diagnosis_strongest_evidence": items_a,
            "against_diagnosis_strongest_evidence": items_b,
            "summary": f"{marker} overall summary",
        }
    return {
        "diagnosisA_strongest_evidence": items_a,
        "diagnosisB_strongest_evidence": items_b,
        "summary": f"{marker} differential summary",
    }


def _response_record(
    request: FeedbackRequest,
    *,
    marker: str = "fake",
    attempt_count: int = 1,
    provider_response_id: str | None = "response_fake",
) -> dict[str, object]:
    return {
        "schema_version": RESPONSE_RECORD_SCHEMA_VERSION,
        "request_id": request.request_id,
        "run_id": request.run_id,
        "prompt_sha256": request.prompt_sha256,
        "config_schema_version": request.config_schema_version,
        "prompt_schema_version": request.prompt_schema_version,
        "response_schema_version": request.response_schema_version,
        "model_profile": request.model_profile,
        "model_id": request.model_id,
        "model_settings": request.model_settings,
        "parsed_payload": _payload_for(request, marker=marker),
        "response_payload_sha256": "written-by-storage",
        "provider_response_id": provider_response_id,
        "attempt_count": attempt_count,
        "started_at": "2026-08-12T00:00:00+00:00",
        "finished_at": "2026-08-12T00:00:01+00:00",
        "latency_seconds": 1.0,
        "token_usage": {"prompt_tokens": 10, "completion_tokens": 20, "unsafe": 999},
        "safe_finish_status": {"finish_reason": "stop", "unsafe": SECRET_SENTINEL},
    }


def _complete_directly(run_dir: Path) -> tuple[FeedbackRequest, ...]:
    requests = read_manifest(run_dir / "manifest.csv")
    ledger = LedgerStore(run_dir, requests)
    for request in requests:
        write_response_record(run_dir, request, _response_record(request))
        write_attempt_record(
            run_dir,
            {
                "request_id": request.request_id,
                "attempt": 1,
                "status": "success",
                "failure_class": "",
                "http_status": None,
                "retry_after_seconds": None,
                "started_at": "2026-08-12T00:00:00+00:00",
                "finished_at": "2026-08-12T00:00:01+00:00",
                "latency_seconds": 1.0,
            },
        )
        ledger.update(
            request.request_id,
            status="success",
            attempt_count=1,
            updated_at="2026-08-12T00:00:01+00:00",
        )
    return requests


def _read_ledger(run_dir: Path) -> list[dict[str, str]]:
    with (run_dir / "run_ledger.csv").open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _artifact_text(run_dir: Path) -> str:
    chunks: list[str] = []
    for path in sorted(item for item in run_dir.rglob("*") if item.is_file()):
        if path.suffix.lower() == ".xlsx":
            continue
        chunks.append(path.read_text(encoding="utf-8", errors="ignore"))
    return "\n".join(chunks)


def _assert_secret_absent_from_artifacts(run_dir: Path) -> None:
    assert SECRET_SENTINEL not in _artifact_text(run_dir)
    for path in run_dir.rglob("*.xlsx"):
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                assert SECRET_SENTINEL.encode() not in archive.read(name)


def test_frozen_config_matches_notebook_diagnoses_categories_and_findings(
    feedback_config: FeedbackConfig,
    contract: dict[str, object],
) -> None:
    assert feedback_config.correct_diagnosis == contract["correct_diagnosis"]
    assert list(feedback_config.diagnoses) == contract["diagnoses"]
    assert len(feedback_config.diagnoses) == 28
    assert len(feedback_config.alternatives) == 27
    assert feedback_config.alternatives[22] == "Rheumatoid Athritis"

    expected_categories = contract["categories"]
    assert isinstance(expected_categories, list)
    assert [category.key for category in feedback_config.categories] == [
        item["key"] for item in expected_categories
    ]
    assert [category.order for category in feedback_config.categories] == list(range(1, 7))
    for actual, expected in zip(feedback_config.categories, expected_categories, strict=True):
        assert (
            hashlib.sha256(actual.description.encode()).hexdigest()
            == expected["description_sha256"]
        )
        assert [
            {"finding": detail.finding, "status": detail.status} for detail in actual.details
        ] == expected["details"]


def test_frozen_prompts_and_inventory_hashes(
    feedback_config: FeedbackConfig,
    contract: dict[str, object],
) -> None:
    inventory = contract["request_inventory"]
    fingerprints = contract["prompt_fingerprints"]
    assert isinstance(inventory, dict)
    assert isinstance(fingerprints, dict)

    for mode in ("only_overall", "all"):
        requests = build_feedback_requests(feedback_config, category_mode=mode)
        expected = inventory[mode]
        assert len(requests) == expected["total"]
        assert Counter(request.surface for request in requests) == expected["by_surface"]
        assert requests[0].request_id == expected["first_request_id"]
        assert requests[-1].request_id == expected["last_request_id"]
        assert requests[0].run_id == expected["base_run_id"]
        assert {request.run_id for request in requests} == {expected["base_run_id"]}
        assert _ordered_messages_sha256(requests) == expected["ordered_messages_sha256"]
        assert [request.request_order for request in requests] == list(range(1, len(requests) + 1))

    sample_requests = build_feedback_requests(feedback_config, category_mode="only_overall")
    for surface in contract["surface_order"]:
        request = next(item for item in sample_requests if item.surface == surface)
        expected = fingerprints[surface]
        assert request.messages[0]["content"] == fingerprints["system"]["text"]
        assert request.messages[1]["content"] == expected["user_prompt"]
        assert (
            hashlib.sha256(request.messages[1]["content"].encode()).hexdigest()
            == expected["user_prompt_sha256"]
        )
        assert prompt_sha256(request.messages) == request.prompt_sha256


def test_prompt_hash_normalizes_only_newline_encoding() -> None:
    lf = [
        {"role": "system", "content": "first\nsecond\n"},
        {"role": "user", "content": "trailing spaces stay  \n"},
    ]
    crlf = [
        {"role": "system", "content": "first\r\nsecond\r\n"},
        {"role": "user", "content": "trailing spaces stay  \r"},
    ]
    assert prompt_sha256(lf) == prompt_sha256(crlf)

    no_trailing_newline = deepcopy(lf)
    no_trailing_newline[-1]["content"] = no_trailing_newline[-1]["content"].rstrip("\n")
    assert prompt_sha256(no_trailing_newline) != prompt_sha256(lf)

    removed_space = deepcopy(lf)
    removed_space[-1]["content"] = "trailing spaces stay\n"
    assert prompt_sha256(removed_space) != prompt_sha256(lf)


def test_frozen_model_profiles_and_workbook_contract(
    feedback_config: FeedbackConfig,
    contract: dict[str, object],
) -> None:
    expected_profiles = contract["model_profiles"]
    assert {
        profile.name: {"model_id": profile.model_id, "settings": profile.settings}
        for profile in feedback_config.model_profiles
    } == {
        name: {"model_id": value["model_id"], "settings": value["settings"]}
        for name, value in expected_profiles.items()
    }
    assert feedback_config.default_model_profile == "o3-mini"
    assert feedback_config.model().model_id == "o3-mini-2025-01-31"
    assert feedback_config.model().settings == {"reasoning_effort": "high"}
    for profile_name, expected in expected_profiles.items():
        requests = build_feedback_requests(
            feedback_config,
            category_mode="only_overall",
            model_profile=profile_name,
        )
        assert requests[0].run_id == expected["default_base_run_id"]

    workbook_contract = contract["legacy_workbook_contract"]
    default_requests = build_feedback_requests(feedback_config, category_mode="only_overall")
    assert len({request.workbook_path for request in default_requests}) == 110
    for surface in contract["surface_order"]:
        actual = [
            request.workbook_path.removeprefix("workbooks/")
            for request in default_requests
            if request.surface == surface
        ]
        assert actual == workbook_contract["filename_order_by_surface"][surface]
    assert {request.sheet_name for request in default_requests} == set(
        workbook_contract["sheet_order_only_overall"]
    )

    expanded = build_feedback_requests(feedback_config, category_mode="all")
    grouped: dict[str, list[FeedbackRequest]] = {}
    for request in expanded:
        grouped.setdefault(request.workbook_path, []).append(request)
    assert len(grouped) == 110
    for requests in grouped.values():
        assert [request.sheet_name for request in requests] == workbook_contract["sheet_order_all"]
        assert [request.category_order for request in requests] == list(range(1, 7))


def test_manifest_columns_freeze_all_identity_and_destination_fields() -> None:
    assert MANIFEST_COLUMNS == (
        "request_order",
        "request_id",
        "run_id",
        "surface",
        "correct_diagnosis",
        "target_diagnosis",
        "comparison_diagnosis",
        "diagnosis_order",
        "category_key",
        "category_order",
        "config_schema_version",
        "prompt_schema_version",
        "response_schema_version",
        "model_profile",
        "model_id",
        "model_settings",
        "messages",
        "prompt_sha256",
        "expected_response_type",
        "response_path",
        "workbook_path",
        "sheet_name",
        "initial_status",
    )


def _duplicate_diagnosis(value: dict[str, object]) -> None:
    diagnoses = value["diagnoses"]
    assert isinstance(diagnoses, list)
    diagnoses[1] = diagnoses[0]


def _duplicate_category(value: dict[str, object]) -> None:
    categories = value["categories"]
    assert isinstance(categories, list)
    categories[1]["key"] = categories[0]["key"]


def _invalid_category_sources(value: dict[str, object]) -> None:
    categories = value["categories"]
    assert isinstance(categories, list)
    categories[0]["detail_sources"] = ["hist"]


def _invalid_finding_status(value: dict[str, object]) -> None:
    value["categories"][0]["details"][0]["status"] = "unknown"


def _duplicate_finding(value: dict[str, object]) -> None:
    value["categories"][0]["details"][1]["finding"] = value["categories"][0]["details"][0][
        "finding"
    ]


def _unknown_top_level_field(value: dict[str, object]) -> None:
    value["unsupported_field"] = True


def _unknown_runtime_field(value: dict[str, object]) -> None:
    value["runtime_defaults"]["unsupported_field"] = True


def _duplicate_surface(value: dict[str, object]) -> None:
    value["enabled_surfaces"].append(value["enabled_surfaces"][0])


def _unsupported_surface(value: dict[str, object]) -> None:
    value["enabled_surfaces"][0] = "overall/imaginary"


def _empty_surfaces(value: dict[str, object]) -> None:
    value["enabled_surfaces"] = []


def _absolute_output(value: dict[str, object]) -> None:
    value["output_root"] = "/tmp/feedback"


def _parent_output(value: dict[str, object]) -> None:
    value["output_root"] = "../feedback"


def _zero_workers(value: dict[str, object]) -> None:
    value["runtime_defaults"]["max_workers"] = 0


def _zero_timeout(value: dict[str, object]) -> None:
    value["runtime_defaults"]["request_timeout_seconds"] = 0


def _zero_attempts(value: dict[str, object]) -> None:
    value["runtime_defaults"]["max_attempts"] = 0


def _zero_initial_retry_delay(value: dict[str, object]) -> None:
    value["runtime_defaults"]["retry_initial_delay_seconds"] = 0


def _zero_max_retry_delay(value: dict[str, object]) -> None:
    value["runtime_defaults"]["retry_max_delay_seconds"] = 0


def _backoff_reversed(value: dict[str, object]) -> None:
    value["runtime_defaults"]["retry_initial_delay_seconds"] = 5
    value["runtime_defaults"]["retry_max_delay_seconds"] = 1


def _unknown_default_model(value: dict[str, object]) -> None:
    value["default_model_profile"] = "missing"


def _unsupported_model_setting(value: dict[str, object]) -> None:
    value["model_profiles"]["o3-mini"]["settings"]["made_up_parameter"] = True


def _excess_retry_jitter(value: dict[str, object]) -> None:
    value["runtime_defaults"]["retry_jitter_fraction"] = 0.26


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        (_duplicate_diagnosis, "duplicates"),
        (_duplicate_category, "duplicate"),
        (_invalid_category_sources, "exactly one"),
        (_invalid_finding_status, "present.*absent"),
        (_duplicate_finding, "duplicate findings"),
        (_unknown_top_level_field, "unsupported field"),
        (_unknown_runtime_field, "unsupported field"),
        (_duplicate_surface, "duplicates"),
        (_unsupported_surface, "unsupported"),
        (_empty_surfaces, "must not be empty"),
        (_absolute_output, "repository-relative"),
        (_parent_output, "repository-relative"),
        (_zero_workers, "positive integer"),
        (_zero_timeout, "positive"),
        (_zero_attempts, "positive integer"),
        (_zero_initial_retry_delay, "positive"),
        (_zero_max_retry_delay, "positive"),
        (_backoff_reversed, "must be >="),
        (_unknown_default_model, "configured model profile"),
        (_unsupported_model_setting, "unsupported.*setting"),
        (_excess_retry_jitter, "jitter.*0.25"),
    ],
)
def test_configuration_rejects_invalid_contracts(
    tmp_path: Path,
    mutation: Callable[[dict[str, object]], None],
    message: str,
) -> None:
    raw = _raw_config()
    mutation(raw)
    with pytest.raises(ValueError, match=message):
        load_feedback_config(_write_raw_config(tmp_path, raw))


def test_configuration_rejects_duplicate_yaml_keys_and_nonfinite_runtime(
    tmp_path: Path,
) -> None:
    duplicate = tmp_path / "duplicate.yaml"
    duplicate.write_text("schema_version: one\nschema_version: two\n", encoding="utf-8")
    with pytest.raises(ValueError, match="duplicate key"):
        load_feedback_config(duplicate)

    raw = CONFIG_PATH.read_text(encoding="utf-8").replace(
        '"request_timeout_seconds": 120.0', '"request_timeout_seconds": .inf'
    )
    nonfinite = tmp_path / "nonfinite.yaml"
    nonfinite.write_text(raw, encoding="utf-8")
    with pytest.raises(ValueError, match="finite"):
        load_feedback_config(nonfinite)


def _mutate_identity_config(raw: dict[str, object], case: str) -> None:
    if case == "correct_diagnosis":
        raw["correct_diagnosis"] = "Changed correct diagnosis"
        raw["diagnoses"][0] = "Changed correct diagnosis"
    elif case == "alternative_diagnosis":
        raw["diagnoses"][1] = "Changed alternative"
    elif case == "category_key":
        raw["categories"][0]["key"] = "changed-hpi"
        aggregate_sources = raw["categories"][-1]["detail_sources"]
        aggregate_sources[aggregate_sources.index("hpi")] = "changed-hpi"
    elif case == "category_description":
        raw["categories"][0]["description"] += " changed"
    elif case == "finding":
        raw["categories"][0]["details"][0]["finding"] += " changed"
    elif case == "finding_status":
        raw["categories"][0]["details"][0]["status"] = "absent"
    elif case == "category_order":
        raw["categories"][0], raw["categories"][1] = (
            raw["categories"][1],
            raw["categories"][0],
        )
    elif case == "surface_order":
        raw["enabled_surfaces"][0], raw["enabled_surfaces"][1] = (
            raw["enabled_surfaces"][1],
            raw["enabled_surfaces"][0],
        )
    elif case == "prompt_schema":
        raw["prompt_schema_version"] += "-changed"
    elif case == "response_schema":
        raw["response_schema_version"] += "-changed"
    elif case == "model_id":
        raw["model_profiles"]["o3-mini"]["model_id"] += "-changed"
    elif case == "model_settings":
        raw["model_profiles"]["o3-mini"]["settings"]["reasoning_effort"] = "medium"
    else:  # pragma: no cover - parameter list is the exhaustive routing table.
        raise AssertionError(case)


@pytest.mark.parametrize(
    "case",
    [
        "correct_diagnosis",
        "alternative_diagnosis",
        "category_key",
        "category_description",
        "finding",
        "finding_status",
        "category_order",
        "surface_order",
        "prompt_schema",
        "response_schema",
        "model_id",
        "model_settings",
    ],
)
def test_every_output_affecting_config_change_changes_request_identity(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
    case: str,
) -> None:
    baseline = build_feedback_requests(feedback_config, category_mode="all")
    raw = _raw_config()
    _mutate_identity_config(raw, case)
    changed_config = load_feedback_config(_write_raw_config(tmp_path, raw))
    changed = build_feedback_requests(changed_config, category_mode="all")

    assert [request.request_id for request in changed] != [
        request.request_id for request in baseline
    ]
    assert changed[0].run_id != baseline[0].run_id


def test_operational_fields_do_not_change_requests_or_base_identity(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    baseline = build_feedback_requests(feedback_config)
    raw = _raw_config()
    raw["output_root"] = "artifacts/another-relative-root"
    raw["runtime_defaults"].update(
        {
            "max_workers": 1,
            "request_timeout_seconds": 1,
            "max_attempts": 9,
            "retry_initial_delay_seconds": 0.5,
            "retry_max_delay_seconds": 4,
            "retry_jitter_fraction": 0,
        }
    )
    changed = build_feedback_requests(load_feedback_config(_write_raw_config(tmp_path, raw)))
    assert [request.request_id for request in changed] == [
        request.request_id for request in baseline
    ]
    assert changed[0].run_id == baseline[0].run_id


def test_config_schema_changes_request_ids_and_run_fingerprint(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    baseline = build_feedback_requests(feedback_config)
    raw = _raw_config()
    raw["schema_version"] += "-changed"
    changed = build_feedback_requests(load_feedback_config(_write_raw_config(tmp_path, raw)))
    assert [request.request_id for request in changed] != [
        request.request_id for request in baseline
    ]
    assert changed[0].run_id != baseline[0].run_id


def test_manifest_creation_is_atomic_idempotent_and_fail_closed(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    runs_root = tmp_path / "runs"
    first = build_or_load_manifest(feedback_config, output_root=runs_root)
    tracked = {path.name: path.read_bytes() for path in first.run_dir.iterdir() if path.is_file()}
    mtimes = {path.name: path.stat().st_mtime_ns for path in first.run_dir.iterdir()}
    second = build_or_load_manifest(feedback_config, output_root=runs_root)
    assert second.run_dir == first.run_dir
    assert [request.request_id for request in second.requests] == [
        request.request_id for request in first.requests
    ]
    assert tracked == {
        path.name: path.read_bytes() for path in first.run_dir.iterdir() if path.is_file()
    }
    assert mtimes == {path.name: path.stat().st_mtime_ns for path in first.run_dir.iterdir()}

    metadata_path = first.run_dir / "run_metadata.json"
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    metadata["run_fingerprint"] = "0" * 64
    metadata_path.write_text(json.dumps(metadata), encoding="utf-8")
    with pytest.raises(FileExistsError, match="fingerprint"):
        build_or_load_manifest(feedback_config, output_root=runs_root)


def test_manifest_tampering_is_detected(tmp_path: Path, feedback_config: FeedbackConfig) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    rows: list[dict[str, str]]
    with build.manifest_path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = reader.fieldnames
        rows = list(reader)
    assert fieldnames is not None
    messages = json.loads(rows[0]["messages"])
    messages[1]["content"] += "tampered"
    rows[0]["messages"] = canonical_json(messages)
    with build.manifest_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    with pytest.raises(ValueError, match="Prompt hash mismatch"):
        read_manifest(build.manifest_path)
    with pytest.raises(FileExistsError, match="manifest bytes"):
        build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")


def test_manifest_partial_write_failure_leaves_no_run(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from dx_chat_entropy.feedback import manifest as manifest_module

    original = manifest_module._write_file
    calls = 0

    def fail_midpoint(path: Path, content: bytes) -> None:
        nonlocal calls
        calls += 1
        if calls == 3:
            raise OSError("synthetic atomic failure")
        original(path, content)

    monkeypatch.setattr(manifest_module, "_write_file", fail_midpoint)
    runs_root = tmp_path / "runs"
    with pytest.raises(OSError, match="synthetic atomic failure"):
        build_or_load_manifest(feedback_config, output_root=runs_root)
    assert not any(path.name.startswith("fb_") for path in runs_root.iterdir())
    assert not any(path.name.startswith(".fb_") for path in runs_root.iterdir())


def test_response_shapes_require_five_complete_items_and_nonempty_summary(
    feedback_config: FeedbackConfig,
) -> None:
    requests = build_feedback_requests(feedback_config)
    differential = next(item for item in requests if item.expected_response_type == "differential")
    for request in (requests[0], differential):
        payload = _payload_for(request)
        assert validate_response_payload(request.expected_response_type, payload) == payload

        list_key = next(key for key, value in payload.items() if isinstance(value, list))
        short = deepcopy(payload)
        short[list_key].pop()
        with pytest.raises(ValueError, match="exactly 5"):
            validate_response_payload(request.expected_response_type, short)

        empty_finding = deepcopy(payload)
        empty_finding[list_key][0]["finding"] = "  "
        with pytest.raises(ValueError, match="non-empty"):
            validate_response_payload(request.expected_response_type, empty_finding)

        empty_summary = deepcopy(payload)
        empty_summary["summary"] = ""
        with pytest.raises(ValueError, match="summary.*non-empty"):
            validate_response_payload(request.expected_response_type, empty_summary)

        extra_top_level = deepcopy(payload)
        extra_top_level["unexpected"] = SECRET_SENTINEL
        with pytest.raises(ValueError, match="exactly the required fields"):
            validate_response_payload(request.expected_response_type, extra_top_level)

        extra_item_field = deepcopy(payload)
        extra_item_field[list_key][0]["unexpected"] = SECRET_SENTINEL
        with pytest.raises(ValueError, match="exactly the required fields"):
            validate_response_payload(request.expected_response_type, extra_item_field)


def test_provider_result_contract_has_provider_independent_aliases() -> None:
    result = ProviderResult(
        payload={"answer": "fake"},
        provider_response_id="response_fake",
        token_usage={"total_tokens": 3},
        finish_reason="stop",
        status_metadata={"status": "completed"},
    )
    assert result.parsed_payload == {"answer": "fake"}
    assert result.usage == {"total_tokens": 3}
    assert result.safe_metadata == {"status": "completed"}


def test_two_public_response_types_round_trip_validated_payloads(
    feedback_config: FeedbackConfig,
) -> None:
    requests = build_feedback_requests(feedback_config)
    overall_request = requests[0]
    differential_request = next(
        request for request in requests if request.expected_response_type == "differential"
    )
    overall_payload = _payload_for(overall_request)
    differential_payload = _payload_for(differential_request)
    assert OverallFeedbackResponse.validate(overall_payload).to_dict() == overall_payload
    assert DifferentialFeedbackResponse.validate(differential_payload).to_dict() == (
        differential_payload
    )


def test_atomic_write_preserves_destination_and_removes_temp_on_replace_failure(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from dx_chat_entropy.feedback import storage as storage_module

    destination = tmp_path / "record.json"
    destination.write_bytes(b"old")

    def fail_replace(_source: object, _destination: object) -> None:
        raise OSError("synthetic replace failure")

    monkeypatch.setattr(storage_module.os, "replace", fail_replace)
    with pytest.raises(OSError, match="synthetic replace failure"):
        atomic_write_bytes(destination, b"new")
    assert destination.read_bytes() == b"old"
    assert not list(tmp_path.glob("*.tmp"))


def test_ledger_update_keeps_memory_and_disk_in_sync_when_atomic_write_fails(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from dx_chat_entropy.feedback import storage as storage_module

    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    ledger = LedgerStore(build.run_dir, build.requests)
    before_snapshot = ledger.snapshot()
    before_disk = (build.run_dir / "run_ledger.csv").read_bytes()

    def fail_csv(*_args: object, **_kwargs: object) -> None:
        raise OSError("synthetic ledger checkpoint failure")

    monkeypatch.setattr(storage_module, "atomic_write_csv", fail_csv)
    with pytest.raises(OSError, match="synthetic ledger checkpoint failure"):
        ledger.update(
            build.requests[0].request_id,
            status="running",
            attempt_count=1,
            updated_at="2026-08-12T00:00:00+00:00",
        )

    assert ledger.snapshot() == before_snapshot
    assert (build.run_dir / "run_ledger.csv").read_bytes() == before_disk


def test_response_record_round_trip_sanitizes_metadata_and_detects_tampering(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    request = build.requests[0]
    unsafe_record = _response_record(request)
    unsafe_record[SECRET_SENTINEL] = SECRET_SENTINEL
    with pytest.raises(ValueError, match="unexpected fields") as extra_error:
        write_response_record(build.run_dir, request, unsafe_record)
    assert SECRET_SENTINEL not in str(extra_error.value)
    assert not (build.run_dir / request.response_path).exists()
    assert SECRET_SENTINEL not in _artifact_text(build.run_dir)

    write_response_record(
        build.run_dir,
        request,
        _response_record(request, provider_response_id=SECRET_SENTINEL),
    )
    path = build.run_dir / request.response_path
    record = validate_response_record(request, path)
    assert record["provider_response_id"] is None
    assert record["config_schema_version"] == request.config_schema_version
    assert record["prompt_schema_version"] == request.prompt_schema_version
    assert record["token_usage"] == {"completion_tokens": 20, "prompt_tokens": 10}
    assert record["safe_finish_status"] == {"finish_reason": "stop"}
    assert SECRET_SENTINEL not in path.read_text(encoding="utf-8")

    raw = json.loads(path.read_text(encoding="utf-8"))
    raw["unexpected"] = "provider detail"
    path.write_text(json.dumps(raw), encoding="utf-8")
    with pytest.raises(ValueError, match="unexpected fields"):
        validate_response_record(request, path)

    write_response_record(build.run_dir, request, _response_record(request))
    raw = json.loads(path.read_text(encoding="utf-8"))
    raw["parsed_payload"]["summary"] = "tampered"
    path.write_text(json.dumps(raw), encoding="utf-8")
    with pytest.raises(ValueError, match="payload hash"):
        validate_response_record(request, path)

    for field in ("config_schema_version", "prompt_schema_version"):
        write_response_record(build.run_dir, request, _response_record(request))
        raw = json.loads(path.read_text(encoding="utf-8"))
        raw[field] = f"{raw[field]}-tampered"
        path.write_text(json.dumps(raw), encoding="utf-8")
        with pytest.raises(ValueError, match=field):
            validate_response_record(request, path)

    escaping = replace(request, response_path="../outside.json")
    with pytest.raises(ValueError, match="safe and relative"):
        write_response_record(build.run_dir, escaping, _response_record(escaping))
    assert not (build.run_dir.parent / "outside.json").exists()


def test_attempt_and_ledger_metadata_are_strictly_allowlisted(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    request = build.requests[0]
    valid_attempt: dict[str, object] = {
        "request_id": request.request_id,
        "attempt": 1,
        "status": "transient_failure",
        "failure_class": "transient_provider",
        "http_status": 429,
        "retry_after_seconds": 7.0,
        "started_at": "2026-08-12T00:00:00+00:00",
        "finished_at": "2026-08-12T00:00:01+00:00",
        "latency_seconds": 1.0,
    }
    invalid_values = (
        ("status", SECRET_SENTINEL),
        ("failure_class", SECRET_SENTINEL),
        ("http_status", 99),
        ("http_status", "429"),
        ("retry_after_seconds", float("inf")),
        ("retry_after_seconds", 3600.1),
        ("started_at", SECRET_SENTINEL),
        ("finished_at", "2026-08-11T23:59:59+00:00"),
        ("latency_seconds", -1),
    )
    for field, value in invalid_values:
        candidate = dict(valid_attempt)
        candidate[field] = value
        with pytest.raises(ValueError) as error:
            write_attempt_record(build.run_dir, candidate)
        assert SECRET_SENTINEL not in str(error.value)
    assert not (build.run_dir / "attempts" / request.request_id).exists()
    assert SECRET_SENTINEL not in _artifact_text(build.run_dir)

    path = write_attempt_record(build.run_dir, valid_attempt)
    stored_attempt = validate_attempt_record(
        path,
        expected_request_id=request.request_id,
        expected_attempt=1,
    )
    assert stored_attempt["http_status"] == 429
    assert stored_attempt["retry_after_seconds"] == 7.0
    with pytest.raises(ValueError, match="expected path"):
        validate_attempt_record(path, expected_attempt=2)
    with pytest.raises(ValueError, match="1 through 9999"):
        validate_attempt_record(path, expected_attempt=10_000)

    maximum_attempt = dict(valid_attempt)
    maximum_attempt["attempt"] = 9999
    maximum_path = write_attempt_record(build.run_dir, maximum_attempt)
    assert maximum_path.name == "9999.json"
    assert validate_attempt_record(maximum_path, expected_attempt=9999)["attempt"] == 9999
    for invalid_attempt in (0, 10_000):
        candidate = dict(valid_attempt)
        candidate["attempt"] = invalid_attempt
        with pytest.raises(ValueError, match="1 through 9999"):
            write_attempt_record(build.run_dir, candidate)

    ledger = LedgerStore(build.run_dir, build.requests)
    for status, failure_class in (
        ("transient_failure", "stale_running"),
        ("invalid", "invalid_stored_response"),
        ("permanent_failure", "attempt_cap"),
        ("transient_failure", "interrupted"),
    ):
        ledger.update(
            request.request_id,
            status=status,
            updated_at="2026-08-12T00:00:02+00:00",
            failure_class=failure_class,
        )

    with pytest.raises(ValueError, match="failure_class is not allowlisted") as failure_error:
        ledger.update(
            request.request_id,
            status="transient_failure",
            updated_at="2026-08-12T00:00:03+00:00",
            failure_class=SECRET_SENTINEL,
        )
    assert SECRET_SENTINEL not in str(failure_error.value)
    with pytest.raises(ValueError, match="http_status"):
        ledger.update(
            request.request_id,
            status="permanent_failure",
            updated_at="2026-08-12T00:00:03+00:00",
            failure_class="permanent_provider",
            http_status=600,
        )
    assert SECRET_SENTINEL not in _artifact_text(build.run_dir)

    ledger_path = build.run_dir / "run_ledger.csv"
    original_ledger = ledger_path.read_text(encoding="utf-8")

    def rewrite_first_ledger_row(**updates: str) -> None:
        rows = _read_ledger(build.run_dir)
        rows[0].update(updates)
        with ledger_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)

    rewrite_first_ledger_row(failure_class=SECRET_SENTINEL)
    with pytest.raises(ValueError, match="failure_class is not allowlisted") as load_error:
        LedgerStore(build.run_dir, build.requests)
    assert SECRET_SENTINEL not in str(load_error.value)
    ledger_path.write_text(original_ledger, encoding="utf-8")

    rewrite_first_ledger_row(http_status="700")
    with pytest.raises(ValueError, match="http_status"):
        LedgerStore(build.run_dir, build.requests)
    ledger_path.write_text(original_ledger, encoding="utf-8")


def test_audit_strictly_validates_ordered_ledger_contract_without_leaking_status(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    _complete_directly(build.run_dir)
    materialize_workbooks(build.run_dir)
    rows = _read_ledger(build.run_dir)
    target_id = build.requests[0].request_id
    target = next(row for row in rows if row["request_id"] == target_id)
    target.update(
        {
            "status": SECRET_SENTINEL,
            "failure_class": SECRET_SENTINEL,
            "http_status": "700",
            "updated_at": "2026-08-12T00:00:01",
            "response_path": "",
        }
    )
    rows[0], rows[1] = rows[1], rows[0]
    ledger_path = build.run_dir / "run_ledger.csv"
    with ledger_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    result = audit_feedback_run(build.run_dir)

    checks = {finding.check for finding in result.findings}
    assert {
        "ledger_order",
        "ledger_status",
        "ledger_failure_class",
        "ledger_http_status",
        "ledger_updated_at",
        "ledger_response_path",
    } <= checks
    persisted_report = (build.run_dir / "invalid_requests.csv").read_text(encoding="utf-8") + (
        build.run_dir / "quality_summary.json"
    ).read_text(encoding="utf-8")
    rendered_result = json.dumps(
        {
            "findings": [finding.row() for finding in result.findings],
            "summary": result.summary,
        }
    )
    assert SECRET_SENTINEL not in persisted_report
    assert SECRET_SENTINEL not in rendered_result
    assert result.summary["ledger_status_counts"]["invalid_status"] == 1


def test_recompute_creates_monotonic_fresh_runs_and_preserves_request_ids(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    base = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    first = create_recompute_run(base.run_dir)
    second = create_recompute_run(base.run_dir)
    third = create_recompute_run(first)
    assert first.name == f"{base.run_dir.name}_r001"
    assert second.name == f"{base.run_dir.name}_r002"
    assert third.name == f"{base.run_dir.name}_r003"

    base_requests = read_manifest(base.manifest_path)
    recomputed = read_manifest(first / "manifest.csv")
    assert [request.request_id for request in recomputed] == [
        request.request_id for request in base_requests
    ]
    assert {request.run_id for request in recomputed} == {first.name}
    assert {row["status"] for row in _read_ledger(first)} == {"pending"}
    assert not list((first / "responses").iterdir())


@pytest.mark.parametrize(
    ("artifact", "mutation"),
    [
        ("manifest_summary.json", "missing"),
        ("manifest_summary.json", "tamper"),
        ("run_metadata.json", "missing"),
        ("run_metadata.json", "tamper"),
        ("run_ledger.csv", "missing"),
        ("run_ledger.csv", "tamper"),
        ("manifest.csv", "tamper"),
    ],
)
def test_recompute_fails_closed_on_corrupt_source_identity_artifacts(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
    artifact: str,
    mutation: str,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    path = build.run_dir / artifact
    if mutation == "missing":
        path.unlink()
    elif artifact == "manifest_summary.json":
        value = json.loads(path.read_text(encoding="utf-8"))
        value["run_fingerprint"] = "0" * 64
        path.write_text(json.dumps(value), encoding="utf-8")
    elif artifact == "run_metadata.json":
        value = json.loads(path.read_text(encoding="utf-8"))
        value["run_fingerprint"] = "0" * 64
        path.write_text(json.dumps(value), encoding="utf-8")
    elif artifact == "run_ledger.csv":
        rows = _read_ledger(build.run_dir)
        rows[0]["request_id"] = rows[1]["request_id"]
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
    else:
        content = path.read_text(encoding="utf-8")
        path.write_text(content.replace("overall/general", "overall/specific", 1), encoding="utf-8")

    with pytest.raises((FileNotFoundError, ValueError)):
        create_recompute_run(build.run_dir)
    adapter = FakeProviderAdapter()
    with pytest.raises((FileNotFoundError, ValueError)):
        execute_feedback_run(
            build.run_dir,
            adapter,
            resume_mode="recompute",
            max_requests=1,
            max_workers=1,
            max_attempts=1,
        )
    assert adapter.call_count == 0
    assert not any(path.name.endswith("_r001") for path in build.run_dir.parent.iterdir())


def test_request_selection_filters_in_manifest_order_and_caps_deterministically(
    feedback_config: FeedbackConfig,
) -> None:
    requests = build_feedback_requests(feedback_config, category_mode="all")
    selected = select_requests(
        tuple(reversed(requests)),
        surfaces=["differential/specific"],
        categories=["test", "subjective-and-historical"],
        diagnoses=["Scleroderma"],
        max_requests=2,
    )
    assert [request.request_order for request in selected] == sorted(
        request.request_order for request in selected
    )
    assert len(selected) == 2
    assert {request.surface for request in selected} == {"differential/specific"}
    assert [request.category_key for request in selected] == [
        "test",
        "subjective-and-historical",
    ]
    assert {request.comparison_diagnosis for request in selected} == {"Scleroderma"}

    exact = select_requests(requests, request_ids=[requests[9].request_id, requests[2].request_id])
    assert exact == [requests[2], requests[9]]
    assert select_requests(requests, max_requests=0) == []
    with pytest.raises(ValueError, match="non-negative"):
        select_requests(requests, max_requests=-1)


def test_dry_run_reconciles_selection_without_provider_calls(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    adapter = FakeProviderAdapter()
    summary = execute_feedback_run(
        build.run_dir,
        adapter,
        max_requests=3,
        dry_run=True,
    )
    assert summary.dry_run
    assert summary.selected_count == 3
    assert summary.pending_count == 3
    assert summary.provider_calls == 0
    assert adapter.call_count == 0
    assert {row["status"] for row in _read_ledger(build.run_dir)} == {"pending"}
    assert not list((build.run_dir / "responses").iterdir())
    assert not list((build.run_dir / "attempts").iterdir())


def test_fake_execution_is_bounded_and_completion_order_independent(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    base = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    serial_dir = create_recompute_run(base.run_dir)
    concurrent_dir = create_recompute_run(base.run_dir)
    selected_ids = [request.request_id for request in base.requests[:12]]

    def result_factory(request: FeedbackRequest) -> ProviderResult:
        # Deliberately finish out of manifest order under concurrency.
        time.sleep(((13 - request.request_order) % 4) * 0.002)
        return ProviderResult(payload=_payload_for(request), finish_reason="stop")

    serial_adapter = FakeProviderAdapter(result_factory=result_factory)
    serial = execute_feedback_run(
        serial_dir,
        serial_adapter,
        request_ids=selected_ids,
        max_workers=1,
        max_attempts=1,
    )
    concurrent_adapter = FakeProviderAdapter(result_factory=result_factory)
    concurrent = execute_feedback_run(
        concurrent_dir,
        concurrent_adapter,
        request_ids=selected_ids,
        max_workers=4,
        max_attempts=1,
    )
    assert serial.success_count == concurrent.success_count == len(selected_ids)
    assert serial.provider_calls == concurrent.provider_calls == len(selected_ids)
    assert serial_adapter.max_active == 1
    assert 2 <= concurrent_adapter.max_active <= 4
    assert set(serial_adapter.calls) == set(concurrent_adapter.calls) == set(selected_ids)

    serial_requests = {
        request.request_id: request for request in read_manifest(serial_dir / "manifest.csv")
    }
    concurrent_requests = {
        request.request_id: request for request in read_manifest(concurrent_dir / "manifest.csv")
    }
    for request_id in selected_ids:
        left = validate_response_record(
            serial_requests[request_id], serial_dir / serial_requests[request_id].response_path
        )
        right = validate_response_record(
            concurrent_requests[request_id],
            concurrent_dir / concurrent_requests[request_id].response_path,
        )
        assert left["parsed_payload"] == right["parsed_payload"]
        assert left["response_payload_sha256"] == right["response_payload_sha256"]


class _HTTPFailure(Exception):
    def __init__(
        self,
        status_code: int,
        *,
        retry_after: str | None = None,
        message: str = "synthetic provider failure",
    ) -> None:
        super().__init__(message)
        self.status_code = status_code
        headers = {} if retry_after is None else {"Retry-After": retry_after}
        self.response = SimpleNamespace(status_code=status_code, headers=headers)


def test_transient_retry_attempt_cap_and_secret_safe_records(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    request = build.requests[0]
    sleeps: list[float] = []
    adapter = FakeProviderAdapter(
        failures={
            request.request_id: [
                TimeoutError(SECRET_SENTINEL),
                ConnectionError(SECRET_SENTINEL),
            ]
        }
    )
    summary = execute_feedback_run(
        build.run_dir,
        adapter,
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=3,
        retry_initial_delay_seconds=1,
        retry_max_delay_seconds=8,
        retry_jitter_fraction=0,
        sleep=sleeps.append,
    )
    assert summary.success_count == 1
    assert summary.provider_calls == 3
    assert adapter.attempts_by_request == {request.request_id: 3}
    assert sleeps == [1, 2]
    attempts = sorted((build.run_dir / "attempts" / request.request_id).glob("*.json"))
    assert len(attempts) == 3
    assert [json.loads(path.read_text())["status"] for path in attempts] == [
        "transient_failure",
        "transient_failure",
        "success",
    ]
    assert _read_ledger(build.run_dir)[0]["attempt_count"] == "3"
    _assert_secret_absent_from_artifacts(build.run_dir)


def test_retry_after_guidance_overrides_backoff(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    request = build.requests[0]
    sleeps: list[float] = []
    adapter = FakeProviderAdapter(
        failures={request.request_id: [_HTTPFailure(429, retry_after="7")]}
    )
    summary = execute_feedback_run(
        build.run_dir,
        adapter,
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=2,
        retry_initial_delay_seconds=1,
        retry_max_delay_seconds=2,
        retry_jitter_fraction=0,
        sleep=sleeps.append,
    )
    assert summary.success_count == 1
    assert adapter.call_count == 2
    assert sleeps == [7]
    first_attempt = json.loads(
        (build.run_dir / "attempts" / request.request_id / "0001.json").read_text()
    )
    assert first_attempt["failure_class"] == "transient_provider"
    assert first_attempt["http_status"] == 429
    assert first_attempt["retry_after_seconds"] == 7


@pytest.mark.parametrize(
    ("failures", "expected_status", "expected_calls"),
    [
        ([_HTTPFailure(401)], "permanent_failure", 1),
        ([ValueError("synthetic refusal")], "invalid", 1),
        ([TimeoutError("one"), TimeoutError("two"), TimeoutError("three")], "transient_failure", 3),
    ],
)
def test_nonretryable_failures_and_transient_attempt_caps(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
    failures: list[BaseException],
    expected_status: str,
    expected_calls: int,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    request = build.requests[0]
    adapter = FakeProviderAdapter(failures={request.request_id: failures})
    summary = execute_feedback_run(
        build.run_dir,
        adapter,
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=3,
        retry_initial_delay_seconds=0,
        retry_max_delay_seconds=0,
        retry_jitter_fraction=0,
        sleep=lambda _delay: None,
    )
    assert adapter.call_count == expected_calls
    assert summary.to_dict()[expected_status] == 1
    row = _read_ledger(build.run_dir)[0]
    assert row["status"] == expected_status
    assert int(row["attempt_count"]) == expected_calls
    attempt_paths = (build.run_dir / "attempts" / request.request_id).glob("*.json")
    assert len(list(attempt_paths)) == expected_calls


def test_arbitrary_failure_class_name_does_not_trigger_retry(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    class PolicyConnectionConfigurationError(Exception):
        pass

    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    request = build.requests[0]
    adapter = FakeProviderAdapter(
        failures={request.request_id: [PolicyConnectionConfigurationError("synthetic")]}
    )

    summary = execute_feedback_run(
        build.run_dir,
        adapter,
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=3,
        retry_initial_delay_seconds=0,
        retry_max_delay_seconds=0,
        retry_jitter_fraction=0,
        sleep=lambda _delay: None,
    )

    assert summary.permanent_failure_count == 1
    assert summary.provider_calls == 1
    assert adapter.call_count == 1
    assert _read_ledger(build.run_dir)[0]["attempt_count"] == "1"
    attempt = validate_attempt_record(
        build.run_dir / "attempts" / request.request_id / "0001.json",
        expected_request_id=request.request_id,
        expected_attempt=1,
    )
    assert attempt["status"] == "permanent_failure"
    assert attempt["failure_class"] == "permanent_provider"


def test_typed_openai_timeout_and_connection_failures_are_retried(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    httpx = pytest.importorskip("httpx")
    openai = pytest.importorskip("openai")
    provider_request = httpx.Request("POST", "https://api.openai.invalid/v1/chat/completions")
    failures = (
        openai.APITimeoutError(request=provider_request),
        openai.APIConnectionError(message="synthetic", request=provider_request),
    )

    for index, failure in enumerate(failures):
        build = build_or_load_manifest(
            feedback_config,
            output_root=tmp_path / f"runs-{index}",
        )
        request = build.requests[0]
        adapter = FakeProviderAdapter(failures={request.request_id: [failure]})

        summary = execute_feedback_run(
            build.run_dir,
            adapter,
            request_ids=[request.request_id],
            max_workers=1,
            max_attempts=2,
            retry_initial_delay_seconds=0,
            retry_max_delay_seconds=0,
            retry_jitter_fraction=0,
            sleep=lambda _delay: None,
        )

        assert summary.success_count == 1
        assert summary.provider_calls == 2
        assert adapter.call_count == 2
        attempt = validate_attempt_record(
            build.run_dir / "attempts" / request.request_id / "0001.json",
            expected_request_id=request.request_id,
            expected_attempt=1,
        )
        assert attempt["status"] == "transient_failure"
        assert attempt["failure_class"] == "transient_provider"


def test_attempt_cap_is_total_across_resume_commands(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    request = build.requests[0]
    failing = FakeProviderAdapter(
        failures={request.request_id: [TimeoutError("one"), TimeoutError("two")]}
    )
    first = execute_feedback_run(
        build.run_dir,
        failing,
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=2,
        retry_initial_delay_seconds=0,
        retry_max_delay_seconds=0,
        retry_jitter_fraction=0,
    )
    assert first.transient_failure_count == 1
    assert failing.call_count == 2

    capped = FakeProviderAdapter()
    second = execute_feedback_run(
        build.run_dir,
        capped,
        resume_mode="repair_invalid",
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=2,
    )
    assert second.transient_failure_count == 1
    assert capped.call_count == 0

    final_adapter = FakeProviderAdapter()
    final = execute_feedback_run(
        build.run_dir,
        final_adapter,
        resume_mode="repair_invalid",
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=3,
    )
    assert final.success_count == 1
    assert final_adapter.call_count == 1
    assert _read_ledger(build.run_dir)[0]["attempt_count"] == "3"


def test_invalid_structured_payload_is_not_retried(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    request = build.requests[0]
    payload = _payload_for(request)
    payload["for_diagnosis_strongest_evidence"].pop()
    adapter = FakeProviderAdapter(results={request.request_id: ProviderResult(payload=payload)})
    summary = execute_feedback_run(
        build.run_dir,
        adapter,
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=3,
    )
    assert summary.invalid_count == 1
    assert adapter.call_count == 1
    assert not (build.run_dir / request.response_path).exists()


def test_independent_failure_does_not_stop_other_requests_and_sets_nonzero_exit(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    selected = build.requests[:4]
    failed = selected[1]
    adapter = FakeProviderAdapter(
        failures={failed.request_id: [_HTTPFailure(401, message=SECRET_SENTINEL)]}
    )
    summary = execute_feedback_run(
        build.run_dir,
        adapter,
        request_ids=[request.request_id for request in selected],
        max_workers=4,
        max_attempts=3,
    )
    assert summary.success_count == 3
    assert summary.permanent_failure_count == 1
    assert summary.provider_calls == 4
    assert summary.failure_count == 1
    assert summary.exit_code == 1
    assert set(adapter.calls) == {request.request_id for request in selected}
    assert SECRET_SENTINEL not in _artifact_text(build.run_dir)


def test_skip_passing_revalidates_tamper_and_reconciles_stale_running(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    first, second, third = build.requests[:3]
    initial = FakeProviderAdapter()
    summary = execute_feedback_run(
        build.run_dir,
        initial,
        request_ids=[first.request_id, second.request_id],
        max_workers=2,
        max_attempts=1,
    )
    assert summary.success_count == 2

    no_call = FakeProviderAdapter()
    skipped = execute_feedback_run(
        build.run_dir,
        no_call,
        request_ids=[first.request_id, second.request_id],
        max_workers=2,
    )
    assert skipped.skipped_count == 2
    assert no_call.call_count == 0

    first_path = build.run_dir / first.response_path
    tampered = json.loads(first_path.read_text())
    tampered["prompt_sha256"] = "0" * 64
    first_path.write_text(json.dumps(tampered), encoding="utf-8")
    repair = FakeProviderAdapter()
    repaired = execute_feedback_run(
        build.run_dir,
        repair,
        request_ids=[first.request_id],
        max_workers=1,
        max_attempts=2,
    )
    assert repaired.success_count == 1
    assert repair.calls == (first.request_id,)
    validate_response_record(first, first_path)

    ledger = LedgerStore(build.run_dir, build.requests)
    ledger.update(
        second.request_id,
        status="running",
        updated_at="2026-08-12T00:00:02+00:00",
    )
    ledger.update(
        third.request_id,
        status="running",
        attempt_count=1,
        updated_at="2026-08-12T00:00:02+00:00",
    )
    stale_adapter = FakeProviderAdapter()
    reconciled = execute_feedback_run(
        build.run_dir,
        stale_adapter,
        request_ids=[second.request_id, third.request_id],
        max_workers=1,
        max_attempts=2,
    )
    assert reconciled.skipped_count == 1
    assert reconciled.success_count == 1
    assert stale_adapter.calls == (third.request_id,)
    rows = {row["request_id"]: row for row in _read_ledger(build.run_dir)}
    assert rows[second.request_id]["status"] == "skipped_existing"
    assert rows[third.request_id]["status"] == "success"


def test_stale_running_without_attempt_record_is_repaired_to_an_auditable_run(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    _complete_directly(build.run_dir)
    request = build.requests[0]
    response_path = build.run_dir / request.response_path
    attempt_one_path = build.run_dir / "attempts" / request.request_id / "0001.json"
    response_path.unlink()
    attempt_one_path.unlink()
    ledger = LedgerStore(build.run_dir, build.requests)
    ledger.update(
        request.request_id,
        status="running",
        attempt_count=1,
        updated_at="2026-08-12T00:00:02+00:00",
    )

    adapter = FakeProviderAdapter()
    summary = execute_feedback_run(
        build.run_dir,
        adapter,
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=2,
    )

    assert summary.success_count == 1
    assert summary.provider_calls == 1
    assert adapter.calls == (request.request_id,)
    interrupted = validate_attempt_record(
        attempt_one_path,
        expected_request_id=request.request_id,
        expected_attempt=1,
    )
    assert interrupted["status"] == "transient_failure"
    assert interrupted["failure_class"] == "interrupted"
    completed = validate_attempt_record(
        build.run_dir / "attempts" / request.request_id / "0002.json",
        expected_request_id=request.request_id,
        expected_attempt=2,
    )
    assert completed["status"] == "success"
    row = {item["request_id"]: item for item in _read_ledger(build.run_dir)}[request.request_id]
    assert row["status"] == "success"
    assert row["attempt_count"] == "2"
    assert validate_response_record(request, response_path)["attempt_count"] == 2

    materialize_workbooks(build.run_dir)
    result = audit_feedback_run(build.run_dir)
    assert result.passes, [finding.row() for finding in result.findings[:10]]


def test_resume_repairs_response_before_attempt_and_failure_before_ledger_crash_windows(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    response_first = build_or_load_manifest(
        feedback_config, output_root=tmp_path / "response-first"
    )
    request = response_first.requests[0]
    write_response_record(response_first.run_dir, request, _response_record(request))
    ledger = LedgerStore(response_first.run_dir, response_first.requests)
    ledger.update(
        request.request_id,
        status="running",
        attempt_count=1,
        updated_at="2026-08-12T00:00:00+00:00",
    )
    adapter = FakeProviderAdapter()
    skipped = execute_feedback_run(
        response_first.run_dir,
        adapter,
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=2,
    )
    assert skipped.skipped_count == 1
    assert adapter.call_count == 0
    synthesized = validate_attempt_record(
        response_first.run_dir / "attempts" / request.request_id / "0001.json",
        expected_request_id=request.request_id,
        expected_attempt=1,
    )
    assert synthesized["status"] == "success"
    response_audit = audit_feedback_run(response_first.run_dir)
    assert not any(
        finding.check.startswith("attempt") and finding.request_id == request.request_id
        for finding in response_audit.findings
    )

    failure_first = build_or_load_manifest(feedback_config, output_root=tmp_path / "failure-first")
    request = failure_first.requests[0]
    write_attempt_record(
        failure_first.run_dir,
        {
            "request_id": request.request_id,
            "attempt": 1,
            "status": "permanent_failure",
            "failure_class": "permanent_provider",
            "http_status": 401,
            "retry_after_seconds": None,
            "started_at": "2026-08-12T00:00:00+00:00",
            "finished_at": "2026-08-12T00:00:01+00:00",
            "latency_seconds": 1.0,
        },
    )
    ledger = LedgerStore(failure_first.run_dir, failure_first.requests)
    ledger.update(
        request.request_id,
        status="running",
        attempt_count=1,
        updated_at="2026-08-12T00:00:00+00:00",
    )
    adapter = FakeProviderAdapter()
    repaired = execute_feedback_run(
        failure_first.run_dir,
        adapter,
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=2,
    )
    assert repaired.success_count == 1
    assert adapter.call_count == 1
    assert (
        validate_attempt_record(
            failure_first.run_dir / "attempts" / request.request_id / "0002.json",
            expected_request_id=request.request_id,
            expected_attempt=2,
        )["status"]
        == "success"
    )
    failure_audit = audit_feedback_run(failure_first.run_dir)
    assert not any(
        finding.check.startswith("attempt") and finding.request_id == request.request_id
        for finding in failure_audit.findings
    )


def test_failed_attempt_checkpoint_write_remains_recoverable_without_inventory_gap(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from dx_chat_entropy.feedback import runtime as runtime_module

    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    request = build.requests[0]
    adapter = FakeProviderAdapter(
        failures={request.request_id: [_HTTPFailure(401, message=SECRET_SENTINEL)]}
    )

    def fail_attempt_checkpoint(*_args: object, **_kwargs: object) -> None:
        raise OSError(f"synthetic attempt checkpoint failure {SECRET_SENTINEL}")

    with monkeypatch.context() as checkpoint_failure:
        checkpoint_failure.setattr(
            runtime_module,
            "write_attempt_record",
            fail_attempt_checkpoint,
        )
        with pytest.raises(OSError, match="synthetic attempt checkpoint failure"):
            execute_feedback_run(
                build.run_dir,
                adapter,
                request_ids=[request.request_id],
                max_workers=1,
                max_attempts=2,
            )

    interrupted_row = {row["request_id"]: row for row in _read_ledger(build.run_dir)}[
        request.request_id
    ]
    assert interrupted_row["status"] == "running"
    assert interrupted_row["attempt_count"] == "1"
    assert not (build.run_dir / "attempts" / request.request_id / "0001.json").exists()

    resumed = execute_feedback_run(
        build.run_dir,
        FakeProviderAdapter(),
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=2,
    )
    assert resumed.success_count == 1
    first_attempt = validate_attempt_record(
        build.run_dir / "attempts" / request.request_id / "0001.json",
        expected_request_id=request.request_id,
        expected_attempt=1,
    )
    assert first_attempt["status"] == "transient_failure"
    assert first_attempt["failure_class"] == "interrupted"
    validate_attempt_record(
        build.run_dir / "attempts" / request.request_id / "0002.json",
        expected_request_id=request.request_id,
        expected_attempt=2,
    )
    assert SECRET_SENTINEL not in _artifact_text(build.run_dir)

    completed = execute_feedback_run(
        build.run_dir,
        FakeProviderAdapter(),
        max_workers=8,
    )
    assert completed.skipped_count == 1
    assert completed.success_count == len(build.requests) - 1
    materialize_workbooks(build.run_dir)
    audit = audit_feedback_run(build.run_dir)
    assert audit.passes, [finding.row() for finding in audit.findings[:10]]


def test_repair_invalid_selects_only_failed_then_applies_stable_cap(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    pending, first, second, third = build.requests[:4]
    ledger = LedgerStore(build.run_dir, build.requests)
    for request, status in (
        (first, "invalid"),
        (second, "permanent_failure"),
        (third, "transient_failure"),
    ):
        ledger.update(
            request.request_id,
            status=status,
            attempt_count=1,
            updated_at="2026-08-12T00:00:00+00:00",
        )

    adapter = FakeProviderAdapter()
    result = execute_feedback_run(
        build.run_dir,
        adapter,
        resume_mode="repair_invalid",
        request_ids=[
            pending.request_id,
            first.request_id,
            second.request_id,
            third.request_id,
        ],
        max_requests=1,
        max_workers=2,
        max_attempts=2,
    )
    assert result.selected_count == 1
    assert result.success_count == 1
    assert adapter.calls == (first.request_id,)
    rows = {row["request_id"]: row for row in _read_ledger(build.run_dir)}
    assert rows[second.request_id]["status"] == "permanent_failure"
    assert rows[third.request_id]["status"] == "transient_failure"
    assert rows[pending.request_id]["status"] == "pending"


def test_runtime_recompute_executes_in_fresh_directory_without_overwrite(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    base = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    adapter = FakeProviderAdapter()
    result = execute_feedback_run(
        base.run_dir,
        adapter,
        resume_mode="recompute",
        max_requests=1,
        max_workers=1,
        max_attempts=1,
    )
    assert result.run_dir.name == f"{base.run_dir.name}_r001"
    assert result.run_dir != base.run_dir
    assert not list((base.run_dir / "responses").iterdir())
    assert len(list((result.run_dir / "responses").glob("*.json"))) == 1
    assert [request.request_id for request in read_manifest(result.run_dir / "manifest.csv")] == [
        request.request_id for request in base.requests
    ]


def test_interruption_stops_submission_and_leaves_recoverable_states(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    selected = build.requests[:6]
    adapter = FakeProviderAdapter(
        failures={selected[0].request_id: [KeyboardInterrupt()]},
        delay_seconds=0.01,
    )
    with pytest.raises(KeyboardInterrupt):
        execute_feedback_run(
            build.run_dir,
            adapter,
            request_ids=[request.request_id for request in selected],
            max_workers=2,
            max_attempts=1,
        )
    assert adapter.call_count <= 2
    rows = {row["request_id"]: row for row in _read_ledger(build.run_dir)}
    assert "running" not in {rows[request.request_id]["status"] for request in selected}
    assert rows[selected[0].request_id]["status"] == "transient_failure"
    assert all(rows[request.request_id]["status"] == "pending" for request in selected[2:])

    repaired = FakeProviderAdapter()
    result = execute_feedback_run(
        build.run_dir,
        repaired,
        resume_mode="repair_invalid",
        request_ids=[selected[0].request_id],
        max_workers=1,
        max_attempts=2,
    )
    assert result.success_count == 1
    assert repaired.calls == (selected[0].request_id,)


@pytest.mark.parametrize(
    ("failed_checkpoint", "response_survives", "resume_provider_calls"),
    [
        ("response", False, 1),
        ("attempt", True, 0),
        ("ledger", True, 0),
    ],
)
def test_atomic_checkpoint_failures_resume_to_an_auditable_run(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
    monkeypatch: pytest.MonkeyPatch,
    failed_checkpoint: str,
    response_survives: bool,
    resume_provider_calls: int,
) -> None:
    from dx_chat_entropy.feedback import storage as storage_module

    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    request = build.requests[0]
    response_path = build.run_dir / request.response_path
    attempt_path = build.run_dir / "attempts" / request.request_id / "0001.json"
    ledger_path = build.run_dir / "run_ledger.csv"
    real_atomic_json = storage_module.atomic_write_json
    real_atomic_csv = storage_module.atomic_write_csv
    failure_injected = False

    def fail_one_json(path: str | Path, payload: object) -> Path:
        nonlocal failure_injected
        destination = Path(path)
        target = response_path if failed_checkpoint == "response" else attempt_path
        if not failure_injected and destination == target:
            failure_injected = True
            raise OSError(f"synthetic {failed_checkpoint} checkpoint failure")
        return real_atomic_json(destination, payload)

    def fail_one_ledger_csv(
        path: str | Path,
        fieldnames: tuple[str, ...] | list[str],
        rows: list[Mapping[str, object]] | tuple[Mapping[str, object], ...],
    ) -> Path:
        nonlocal failure_injected
        target_success = any(
            row.get("request_id") == request.request_id and row.get("status") == "success"
            for row in rows
        )
        if not failure_injected and Path(path) == ledger_path and target_success:
            failure_injected = True
            raise OSError("synthetic ledger checkpoint failure")
        return real_atomic_csv(path, fieldnames, rows)

    if failed_checkpoint == "ledger":
        monkeypatch.setattr(storage_module, "atomic_write_csv", fail_one_ledger_csv)
    else:
        monkeypatch.setattr(storage_module, "atomic_write_json", fail_one_json)

    adapter = FakeProviderAdapter()
    with pytest.raises(OSError, match=f"synthetic {failed_checkpoint} checkpoint failure"):
        execute_feedback_run(
            build.run_dir,
            adapter,
            request_ids=[request.request_id],
            max_workers=1,
            max_attempts=2,
        )

    assert failure_injected
    assert adapter.calls == (request.request_id,)
    assert response_path.exists() is response_survives
    attempt = validate_attempt_record(
        attempt_path,
        expected_request_id=request.request_id,
        expected_attempt=1,
    )
    assert attempt["status"] == ("success" if response_survives else "transient_failure")
    if not response_survives:
        assert attempt["failure_class"] == "interrupted"

    resume_adapter = FakeProviderAdapter()
    resumed = execute_feedback_run(
        build.run_dir,
        resume_adapter,
        request_ids=[request.request_id],
        max_workers=1,
        max_attempts=2,
    )
    assert resume_adapter.call_count == resume_provider_calls
    if response_survives:
        assert resumed.skipped_count == 1
    else:
        assert resumed.success_count == 1
    validate_response_record(request, response_path)

    completion_adapter = FakeProviderAdapter()
    completed = execute_feedback_run(
        build.run_dir,
        completion_adapter,
        max_workers=8,
        max_attempts=2,
    )
    assert completed.failure_count == 0
    assert request.request_id not in completion_adapter.calls
    materialize_workbooks(build.run_dir)
    audit = audit_feedback_run(build.run_dir)
    assert audit.passes, [finding.row() for finding in audit.findings[:10]]
    assert not list(build.run_dir.rglob("*.tmp"))


def test_openai_adapter_uses_injected_client_and_structured_parse_only(
    feedback_config: FeedbackConfig,
) -> None:
    request = build_feedback_requests(feedback_config)[0]

    class Parsed:
        def model_dump(self) -> dict[str, object]:
            return _payload_for(request)

    calls: list[dict[str, object]] = []

    def parse(**kwargs: object) -> object:
        calls.append(dict(kwargs))
        return SimpleNamespace(
            id="chatcmpl-fake",
            choices=[
                SimpleNamespace(
                    message=SimpleNamespace(parsed=Parsed()),
                    finish_reason="stop",
                )
            ],
            usage=SimpleNamespace(model_dump=lambda: {"prompt_tokens": 1, "completion_tokens": 2}),
        )

    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(parse=parse)))
    adapter = OpenAIProviderAdapter(client=client)
    result = adapter.execute(request)
    assert adapter.client is client
    assert len(calls) == 1
    assert calls[0]["model"] == request.model_id
    assert calls[0]["messages"] == request.messages
    assert calls[0]["reasoning_effort"] == "high"
    assert "response_format" in calls[0]
    assert result.payload == _payload_for(request)
    assert result.provider_response_id == "chatcmpl-fake"


def test_direct_cli_real_run_requires_explicit_paid_confirmation(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    namespace = runpy.run_path(str(REPO_ROOT / "scripts" / "run_feedback_pipeline.py"))
    main = namespace["main"]
    monkeypatch.setenv("OPENAI_API_KEY", SECRET_SENTINEL)
    monkeypatch.delenv("CONFIRM_PAID_RUN", raising=False)

    exit_code = main(
        [
            "--repo-root",
            str(REPO_ROOT),
            "run",
            "--runs-root",
            str(tmp_path / "runs"),
        ]
    )
    captured = capsys.readouterr()
    assert exit_code == 2
    assert "CONFIRM_PAID_RUN=1" in captured.err
    assert SECRET_SENTINEL not in captured.err


def test_smoke_command_proves_bounded_concurrency_resume_and_atomic_latest_pointer(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    namespace = runpy.run_path(str(REPO_ROOT / "scripts" / "run_feedback_pipeline.py"))
    main = namespace["main"]
    runs_root = tmp_path / "runs"
    exit_code = main(
        [
            "--repo-root",
            str(REPO_ROOT),
            "smoke",
            "--runs-root",
            str(runs_root),
            "--latency-seconds",
            "0.005",
        ]
    )
    capsys.readouterr()
    assert exit_code == 0
    pointer_path = runs_root / "latest_smoke.json"
    pointer = json.loads(pointer_path.read_text(encoding="utf-8"))
    smoke_path = runs_root / pointer["smoke_summary"]
    smoke = json.loads(smoke_path.read_text(encoding="utf-8"))
    assert smoke["request_count"] == 110
    assert smoke["serial_peak_concurrency"] == 1
    assert 2 <= smoke["four_worker_peak_concurrency"] <= 4
    assert smoke["bounded_concurrency_proved"] is True
    assert smoke["midpoint_failure_first_pass"]["invalid"] == 1
    assert smoke["midpoint_repair"]["success"] == 1
    assert smoke["midpoint_repair"]["failure_count"] == 0
    assert not list(runs_root.rglob("*.tmp"))


def test_full_materialization_is_legacy_compatible_deterministic_and_auditable(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
    contract: dict[str, object],
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    _complete_directly(build.run_dir)
    first = materialize_workbooks(build.run_dir)
    assert first["partial"] is False
    assert first["validated_request_count"] == 110
    assert len(first["workbooks"]) == 110
    expected_paths = {
        f"workbooks/{path}"
        for paths in contract["legacy_workbook_contract"]["filename_order_by_surface"].values()
        for path in paths
    }
    assert {entry["workbook_path"] for entry in first["workbooks"]} == expected_paths
    first_hashes = {entry["workbook_path"]: entry["file_sha256"] for entry in first["workbooks"]}

    sample_entry = first["workbooks"][0]
    sample_path = build.run_dir / sample_entry["workbook_path"]
    workbook = load_workbook(sample_path, read_only=True, data_only=False)
    assert workbook.sheetnames == ["subjective-and-historical"]
    worksheet = workbook[workbook.sheetnames[0]]
    assert worksheet.max_row == 6
    assert worksheet.max_column == 4
    expected_overall_headers = [
        value.format(diagnosis=build.requests[0].target_diagnosis)
        for value in contract["legacy_workbook_contract"]["overall"]["columns"]
    ]
    actual_overall_rows = list(worksheet.iter_rows(values_only=True))
    assert list(actual_overall_rows[0]) == expected_overall_headers
    assert [row[0] for row in actual_overall_rows[1:]] == [
        f"fake A finding {index}" for index in range(1, 6)
    ]
    assert workbook.properties.creator == "dx_chat_entropy"
    workbook.close()
    with zipfile.ZipFile(sample_path) as archive:
        infos = archive.infolist()
        assert [info.filename for info in infos] == sorted(info.filename for info in infos)
        assert {info.date_time for info in infos} == {(1980, 1, 1, 0, 0, 0)}

    differential_request = next(
        request for request in build.requests if request.surface == "differential/general"
    )
    differential_entry = next(
        entry
        for entry in first["workbooks"]
        if entry["workbook_path"] == differential_request.workbook_path
    )
    differential_workbook = load_workbook(
        build.run_dir / differential_entry["workbook_path"], read_only=True
    )
    differential_sheet = differential_workbook[differential_request.sheet_name]
    expected_differential_headers = [
        value.format(
            correct_diagnosis=differential_request.correct_diagnosis,
            comparison_diagnosis=differential_request.comparison_diagnosis,
        )
        for value in contract["legacy_workbook_contract"]["differential"]["columns"]
    ]
    assert list(next(differential_sheet.iter_rows(values_only=True))) == (
        expected_differential_headers
    )
    differential_workbook.close()

    # Cross an OpenXML wall-clock second so openpyxl's implicit modified time
    # cannot accidentally make a nondeterministic implementation appear stable.
    time.sleep(1.05)
    second = materialize_workbooks(build.run_dir)
    second_hashes = {entry["workbook_path"]: entry["file_sha256"] for entry in second["workbooks"]}
    assert second_hashes == first_hashes
    result = audit_feedback_run(build.run_dir)
    assert result.passes, [finding.row() for finding in result.findings[:10]]
    assert result.summary["workbook_count"] == 110
    assert result.summary["validated_response_count"] == 110
    _assert_secret_absent_from_artifacts(build.run_dir)


def test_full_materialization_preflight_writes_nothing_when_responses_missing(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    with pytest.raises(WorkbookMaterializationError, match="missing or invalid"):
        materialize_workbooks(build.run_dir)
    assert not list((build.run_dir / "workbooks").rglob("*.xlsx"))
    assert not (build.run_dir / "workbook_manifest.json").exists()


def test_partial_materialization_writes_only_usable_workbooks_with_visible_label(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(
        feedback_config,
        output_root=tmp_path / "runs",
        category_mode="all",
    )
    first_workbook = build.requests[0].workbook_path
    group = [request for request in build.requests if request.workbook_path == first_workbook]
    assert len(group) == 6
    for request in group[:-1]:
        write_response_record(build.run_dir, request, _response_record(request))

    manifest = materialize_workbooks(build.run_dir, allow_partial=True)
    assert manifest["partial"] is True
    assert manifest["validated_request_count"] == 5
    assert len(manifest["workbooks"]) == 1
    entry = manifest["workbooks"][0]
    assert entry["partial"] is True
    assert entry["workbook_path"].endswith("_PARTIAL.xlsx")
    assert entry["missing_or_invalid_request_ids"] == [group[-1].request_id]
    workbook = load_workbook(build.run_dir / entry["workbook_path"], read_only=True)
    assert workbook.sheetnames == [request.sheet_name for request in group[:-1]] + [
        INCOMPLETE_SHEET_NAME
    ]
    incomplete = list(workbook[INCOMPLETE_SHEET_NAME].iter_rows(values_only=True))
    assert incomplete[1][0] == group[-1].request_id
    workbook.close()


def test_workbook_atomic_replace_failure_leaves_no_partial_artifact(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from dx_chat_entropy.feedback import workbooks as workbook_module

    request = replace(build_feedback_requests(feedback_config)[0], run_id="fb_test_atomic")
    run_dir = tmp_path / request.run_id
    run_dir.mkdir()
    write_response_record(run_dir, request, _response_record(request))
    monkeypatch.setattr(workbook_module, "read_manifest", lambda _path: (request,))

    def fail_replace(_source: object, _destination: object) -> None:
        raise OSError("synthetic workbook replace failure")

    monkeypatch.setattr(workbook_module.os, "replace", fail_replace)
    with pytest.raises(OSError, match="synthetic workbook replace failure"):
        materialize_workbooks(run_dir)
    assert not list(run_dir.rglob("*.xlsx"))
    assert not list(run_dir.rglob("*.tmp"))


def test_audit_detects_corrupt_response_workbook_temp_and_secret_artifacts(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    _complete_directly(build.run_dir)
    manifest = materialize_workbooks(build.run_dir)
    request = build.requests[0]
    response_path = build.run_dir / request.response_path
    response = json.loads(response_path.read_text())
    response["parsed_payload"]["summary"] = "corrupt without rehash"
    response_path.write_text(json.dumps(response), encoding="utf-8")
    duplicate_path = build.run_dir / "responses" / "duplicate.json"
    duplicate_path.write_text(json.dumps(response), encoding="utf-8")
    (build.run_dir / "responses" / "unexpected.txt").write_text(
        "not a response record", encoding="utf-8"
    )

    # Corrupt a different workbook so its valid response identity still requires
    # audit to compare the declared complete-workbook hash.
    workbook_path = build.run_dir / manifest["workbooks"][1]["workbook_path"]
    with workbook_path.open("ab") as handle:
        handle.write(b"corruption")
    (build.run_dir / "stale.partial").write_text("unfinished", encoding="utf-8")
    (build.run_dir / "bad.log").write_text(SECRET_SENTINEL, encoding="utf-8")

    result = audit_feedback_run(build.run_dir)
    checks = {finding.check for finding in result.findings}
    assert not result.passes
    assert "invalid_response" in checks
    assert "duplicate_response" in checks
    assert "unexpected_response_file" in checks
    assert "ledger_response_mismatch" in checks
    assert "workbook_hash" in checks
    assert "temporary_file" in checks
    assert "secret_pattern" in checks
    invalid_text = (build.run_dir / "invalid_requests.csv").read_text(encoding="utf-8")
    quality_text = (build.run_dir / "quality_summary.json").read_text(encoding="utf-8")
    assert SECRET_SENTINEL not in invalid_text
    assert SECRET_SENTINEL not in quality_text


def test_audit_rejects_an_empty_manifest_inventory(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    build.manifest_path.write_text(",".join(MANIFEST_COLUMNS) + "\n", encoding="utf-8")

    result = audit_feedback_run(build.run_dir)

    assert not result.passes
    assert any(finding.check == "manifest_validation" for finding in result.findings)
    assert result.summary["manifest_request_count"] == 0


def test_audit_binds_summary_metadata_fingerprint_and_run_directory_to_manifest(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")

    summary_path = build.run_dir / "manifest_summary.json"
    original_summary = summary_path.read_bytes()
    summary_path.unlink()
    missing_summary = audit_feedback_run(build.run_dir)
    assert any(finding.check == "manifest_summary_missing" for finding in missing_summary.findings)

    summary_path.write_bytes(original_summary)
    summary = json.loads(original_summary)
    summary["total_requests"] += 1
    summary_path.write_text(json.dumps(summary), encoding="utf-8")
    tampered_summary = audit_feedback_run(build.run_dir)
    assert any(
        finding.check == "manifest_summary_identity" for finding in tampered_summary.findings
    )
    summary_path.write_bytes(original_summary)

    metadata_path = build.run_dir / "run_metadata.json"
    original_metadata = metadata_path.read_bytes()
    metadata = json.loads(original_metadata)
    metadata["config_schema_version"] = "feedback-config-tampered"
    fingerprint_payload = {
        "config_schema_version": metadata["config_schema_version"],
        "prompt_schema_version": metadata["prompt_schema_version"],
        "response_schema_version": metadata["response_schema_version"],
        "category_mode": metadata["category_mode"],
        "ordered_request_ids": metadata["ordered_request_ids"],
    }
    metadata["run_fingerprint"] = hashlib.sha256(
        canonical_json(fingerprint_payload).encode("utf-8")
    ).hexdigest()
    metadata_path.write_text(json.dumps(metadata), encoding="utf-8")
    self_consistent_metadata_tamper = audit_feedback_run(build.run_dir)
    assert any(
        finding.check == "run_metadata_identity"
        for finding in self_consistent_metadata_tamper.findings
    )

    metadata_path.write_bytes(original_metadata)
    renamed = build.run_dir.with_name("renamed-feedback-run")
    build.run_dir.rename(renamed)
    renamed_run = audit_feedback_run(renamed)
    assert any(finding.check == "run_identity" for finding in renamed_run.findings)


def test_audit_validates_attempt_schema_identity_inventory_and_counts(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    _complete_directly(build.run_dir)
    materialize_workbooks(build.run_dir)
    baseline = audit_feedback_run(build.run_dir)
    assert baseline.passes, [finding.row() for finding in baseline.findings[:10]]

    request, other_request = build.requests[:2]
    attempt_path = build.run_dir / "attempts" / request.request_id / "0001.json"
    relative_attempt_path = attempt_path.relative_to(build.run_dir).as_posix()
    original_attempt = attempt_path.read_bytes()

    def assert_attempt_finding(result: object, *, path_prefix: str, request_id: str = "") -> None:
        findings = result.findings  # type: ignore[attr-defined]
        assert not result.passes  # type: ignore[attr-defined]
        assert any(
            finding.check.startswith("attempt")
            and finding.path.startswith(path_prefix)
            and (not request_id or finding.request_id == request_id)
            for finding in findings
        ), [finding.row() for finding in findings[:10]]

    attempt_path.write_text("{", encoding="utf-8")
    assert_attempt_finding(audit_feedback_run(build.run_dir), path_prefix=relative_attempt_path)

    attempt_path.write_text(json.dumps({"totally": "wrong"}), encoding="utf-8")
    assert_attempt_finding(audit_feedback_run(build.run_dir), path_prefix=relative_attempt_path)

    attempt_path.write_bytes(original_attempt)
    mismatched = json.loads(original_attempt)
    mismatched["request_id"] = other_request.request_id
    attempt_path.write_text(json.dumps(mismatched), encoding="utf-8")
    assert_attempt_finding(
        audit_feedback_run(build.run_dir),
        path_prefix=relative_attempt_path,
    )

    attempt_path.write_bytes(original_attempt)
    failed_attempt = json.loads(original_attempt)
    failed_attempt.update(
        {
            "status": "permanent_failure",
            "failure_class": "permanent_provider",
            "http_status": 401,
        }
    )
    attempt_path.write_text(json.dumps(failed_attempt), encoding="utf-8")
    inconsistent_status = audit_feedback_run(build.run_dir)
    assert not inconsistent_status.passes
    assert any(
        finding.check == "attempt_status"
        and finding.request_id == request.request_id
        and finding.path == relative_attempt_path
        for finding in inconsistent_status.findings
    )

    attempt_path.write_bytes(original_attempt)
    attempt_path.unlink()
    missing_attempt = audit_feedback_run(build.run_dir)
    assert_attempt_finding(
        missing_attempt,
        path_prefix=f"attempts/{request.request_id}",
        request_id=request.request_id,
    )

    attempt_path.write_bytes(original_attempt)
    extra_path = attempt_path.with_name("0002.json")
    extra_attempt = json.loads(original_attempt)
    extra_attempt["attempt"] = 2
    extra_path.write_text(json.dumps(extra_attempt), encoding="utf-8")
    assert_attempt_finding(
        audit_feedback_run(build.run_dir),
        path_prefix=f"attempts/{request.request_id}",
        request_id=request.request_id,
    )
    extra_path.unlink()

    unexpected_path = build.run_dir / "attempts" / "fr_unexpected" / "0001.json"
    unexpected_path.parent.mkdir()
    unexpected_attempt = json.loads(original_attempt)
    unexpected_attempt["request_id"] = "fr_unexpected"
    unexpected_path.write_text(json.dumps(unexpected_attempt), encoding="utf-8")
    assert_attempt_finding(
        audit_feedback_run(build.run_dir),
        path_prefix=unexpected_path.relative_to(build.run_dir).as_posix(),
    )
    unexpected_path.unlink()

    response_path = build.run_dir / request.response_path
    original_response = response_path.read_bytes()
    response = json.loads(original_response)
    response["attempt_count"] = 2
    response_path.write_text(json.dumps(response), encoding="utf-8")
    mismatch = audit_feedback_run(build.run_dir)
    assert not mismatch.passes
    assert any(
        finding.check.startswith("attempt") and finding.request_id == request.request_id
        for finding in mismatch.findings
    )
    response_path.write_bytes(original_response)
    assert audit_feedback_run(build.run_dir).passes


def test_audit_rejects_multiple_successes_and_every_attempt_after_success(
    tmp_path: Path,
    feedback_config: FeedbackConfig,
) -> None:
    build = build_or_load_manifest(feedback_config, output_root=tmp_path / "runs")
    _complete_directly(build.run_dir)
    request = build.requests[0]
    base_attempt = json.loads(
        (build.run_dir / "attempts" / request.request_id / "0001.json").read_text(encoding="utf-8")
    )
    second_success = dict(base_attempt, attempt=2)
    write_attempt_record(build.run_dir, second_success)
    later_failure = dict(
        base_attempt,
        attempt=3,
        status="transient_failure",
        failure_class="transient_provider",
        http_status=429,
    )
    write_attempt_record(build.run_dir, later_failure)

    result = audit_feedback_run(build.run_dir)

    history = [
        finding
        for finding in result.findings
        if finding.check == "attempt_history" and finding.request_id == request.request_id
    ]
    assert any("more than one successful attempt" in finding.reason for finding in history)
    continued_paths = {
        finding.path
        for finding in history
        if "continues after a successful attempt" in finding.reason
    }
    assert continued_paths == {
        f"attempts/{request.request_id}/0002.json",
        f"attempts/{request.request_id}/0003.json",
    }
